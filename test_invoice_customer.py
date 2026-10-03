import copy
import html
import json
import os
import tempfile
import unittest
from io import BytesIO
from unittest.mock import Mock, patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import automation_pipeline as pipeline
from invoice_generator import build_invoice_xlsx, customer_validation_error
from offer_panel_safe import install_offer_panel


PAYER = {"name": 'ООО «СТРОЙЛОГИСТИКА»', "inn": "7716997861", "kpp": "770801001"}
ITEM = {"position": 1, "order_item_id": 10, "decision": "approved", "requested_name": "Товар",
        "selected": {"name": "Товар"}, "quantity": 2, "unit": "шт", "proposed_unit_price": 100}


class InvoiceCustomerTests(unittest.TestCase):
    def test_missing_incomplete_and_placeholder_payers_cannot_generate_invoice(self):
        for customer in (None, {}, {"name": " "}, {"name": "Заказчик по заявке Закупай", **{k: v for k, v in PAYER.items() if k != "name"}},
                         {"name": PAYER["name"], "inn": PAYER["inn"]}, {**PAYER, "inn": "bad"}, {**PAYER, "kpp": "123"}):
            with self.subTest(customer=customer), self.assertRaisesRegex(ValueError, "плательщика|КПП"):
                build_invoice_xlsx({"items": [ITEM], "customer": customer})

    def test_invoice_prints_actual_payer(self):
        wb = load_workbook(BytesIO(build_invoice_xlsx({"items": [ITEM], "customer": PAYER, "invoice_number": 262})))
        buyer = wb.active["A8"].value
        for value in PAYER.values():
            self.assertIn(value, buyer)
        self.assertNotIn("Заказчик по заявке", buyer)

    def test_individual_entrepreneur_does_not_need_kpp(self):
        self.assertIsNone(customer_validation_error({"name": "ИП Тестовый", "inn": "123456789012"}))

    def test_review_notification_keeps_matching_results_without_bad_attachment(self):
        result = {"items": [ITEM], "order_id": 1, "invoice_number": 262, "summary": {"auto_ready": 1}}
        notification = pipeline._gmail_draft_payload(result, 1)
        self.assertNotIn("attachment_base64", notification)
        self.assertIn("счёт не сформирован", notification["subject"])
        self.assertIn("ИНН плательщика", notification["body"])
        self.assertEqual(result["items"], [ITEM])


class CustomerWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db_patch = patch.multiple(pipeline, DB_PATH=os.path.join(self.tmp.name, "jobs.db"), DATABASE_URL="", _initialized_database_key=None)
        self.db_patch.start()
        self.order = {"id": 37238573, "orderItems": [{"id": 10, "goodName": "Товар", "count": 2, "unit": {"name": "шт"}}]}
        self.result = {"order_id": self.order["id"], "invoice_number": 262, "items": [copy.deepcopy(ITEM)],
                       "summary": {"positions": 1, "auto_ready": 0, "approved": 1, "included_in_invoice": 1}, "status": "ready_for_review"}
        with pipeline._connect() as conn:
            cursor = conn.execute(
                "INSERT INTO automation_jobs (dedupe_key,order_id,event_type,status,order_json,result_json,created_at,updated_at,invoice_number) VALUES (?,?,?,?,?,?,?,?,?)",
                ("test", self.order["id"], "api_order", "ready_for_review", json.dumps(self.order), json.dumps(self.result), "now", "now", 262),
            )
            self.job_id = cursor.lastrowid
        self.fetch = Mock(return_value=self.order)
        self.app = FastAPI()
        pipeline.install_automation_pipeline(self.app, self.fetch)
        install_offer_panel(self.app, Mock(), lambda: {}, "https://example.invalid", lambda x: html.escape(str(x or ""), quote=True),
                            fetch_order_by_id=self.fetch, load_offer_context=pipeline.load_automation_offer_context,
                            build_invoice=build_invoice_xlsx, order_hash=pipeline.commercial_order_hash)
        self.client = TestClient(self.app)
        self.customer_url = f"/dashboard/automation/jobs/{self.job_id}/customer"
        self.invoice_url = f"/dashboard/automation/jobs/{self.job_id}/invoice.xlsx"
        self.customer_form = {"confirm_customer": "CONFIRMED", "customer_name": PAYER["name"], "customer_inn": PAYER["inn"], "customer_kpp": PAYER["kpp"]}

    def tearDown(self):
        self.client.close()
        self.db_patch.stop()
        self.tmp.cleanup()

    def test_old_job_download_is_blocked_until_confirmed_payer_is_saved(self):
        self.assertEqual(self.client.get(self.invoice_url).status_code, 409)
        review = self.client.get(f"/dashboard/automation/jobs/{self.job_id}/review")
        self.assertIn("Реквизиты плательщика", review.text)
        self.assertNotIn("Скачать сформированный счёт", review.text)
        response = self.client.post(self.customer_url, data=self.customer_form, follow_redirects=False)
        self.assertEqual(response.status_code, 303)
        response = self.client.get(self.invoice_url)
        self.assertEqual(response.status_code, 200)
        buyer = load_workbook(BytesIO(response.content)).active["A8"].value
        self.assertIn(PAYER["inn"], buyer)
        context = pipeline.load_automation_offer_context(self.order["id"])
        self.assertNotIn("customer", context["order"])  # Preserve the original request snapshot.
        self.assertEqual(context["result"]["customer"]["kpp"], PAYER["kpp"])

    def test_saving_missing_kpp_or_unconfirmed_payer_is_rejected(self):
        self.assertEqual(self.client.post(self.customer_url, data={**self.customer_form, "customer_kpp": ""}).status_code, 409)
        self.assertEqual(self.client.post(self.customer_url, data={**self.customer_form, "confirm_customer": ""}).status_code, 400)
        self.assertEqual(self.client.get(self.invoice_url).status_code, 409)

    @patch("offer_panel_safe.requests.post")
    def test_direct_submit_without_payer_never_calls_zakupay(self, post):
        response = self.client.post(f"/dashboard/order/{self.order['id']}/offer/submit", data={"confirm_send": "SEND", "producer_offer_number": "262"})
        self.assertEqual(response.status_code, 409)
        self.assertIn("ИНН плательщика", response.json()["detail"])
        post.assert_not_called()
        self.fetch.assert_not_called()

    @patch("offer_panel_safe.requests.get")
    def test_offer_page_blocks_send_and_does_not_suggest_repairing_existing_ids(self, get):
        response = self.client.get(f"/dashboard/order/{self.order['id']}/offer")
        self.assertEqual(response.status_code, 200)
        self.assertIn("type='submit' disabled", response.text)
        self.assertIn("ИНН плательщика", response.text)
        self.assertNotIn("Получить ID позиций", response.text)
        get.assert_not_called()

    @patch("offer_panel_safe.requests.post")
    def test_submit_blocks_different_payer_returned_by_live_order(self, post):
        self.client.post(self.customer_url, data=self.customer_form, follow_redirects=False)
        self.fetch.return_value = {**self.order, "customer": {**PAYER, "inn": "1234567890"}}
        response = self.client.post(f"/dashboard/order/{self.order['id']}/offer/submit", data={"confirm_send": "SEND", "producer_offer_number": "262"})
        self.assertEqual(response.status_code, 409)
        self.assertIn("отличаются", response.json()["detail"])
        post.assert_not_called()

    def test_payer_survives_supplier_search_refresh(self):
        self.client.post(self.customer_url, data=self.customer_form, follow_redirects=False)
        with patch.object(pipeline, "build_vi_draft", return_value=copy.deepcopy(self.result)):
            response = self.client.post(f"/dashboard/automation/jobs/{self.job_id}/refresh", follow_redirects=False)
        self.assertEqual(response.status_code, 303)
        self.assertEqual(self.client.get(self.invoice_url).status_code, 200)
        result = pipeline.load_automation_offer_context(self.order["id"])["result"]
        self.assertEqual(result["customer"]["inn"], PAYER["inn"])
        self.assertEqual(len(result["invoice_customer_history"]), 1)

    def test_payer_cannot_be_changed_while_search_is_running(self):
        with pipeline._connect() as conn:
            conn.execute("UPDATE automation_jobs SET status='processing' WHERE id=?", (self.job_id,))
        self.assertEqual(self.client.post(self.customer_url, data=self.customer_form).status_code, 409)

    def test_payer_cannot_be_changed_during_live_submission(self):
        pipeline.begin_offer_submission(self.job_id, self.order["id"], "attempt", "guid", "snapshot", "payload", "262", "file", {})
        self.assertEqual(self.client.post(self.customer_url, data=self.customer_form).status_code, 409)
        self.assertEqual(self.client.get(self.invoice_url).status_code, 409)

    def test_correcting_sent_invoice_keeps_duplicate_submission_protection(self):
        pipeline.mark_automation_offer_created(self.job_id, offer_id=340691441)
        self.assertEqual(self.client.post(self.customer_url, data=self.customer_form, follow_redirects=False).status_code, 303)
        self.assertEqual(self.client.get(self.invoice_url).status_code, 200)
        result = pipeline.load_automation_offer_context(self.order["id"])["result"]
        self.assertTrue(result["live_offer_created"])
        self.assertEqual(result["live_offer_id"], 340691441)
        response = self.client.post(f"/dashboard/order/{self.order['id']}/offer/submit", data={"confirm_send": "SEND", "producer_offer_number": "262"})
        self.assertEqual(response.status_code, 409)
        self.fetch.assert_not_called()

    def test_list_shows_source_date_and_customer_without_external_lookup(self):
        order = {**self.order, "creationDate": "2026-09-28T21:30:00Z", "customer": {"fullName": 'ООО «Заказчик <источник>»'}}
        with pipeline._connect() as conn:
            conn.execute("UPDATE automation_jobs SET order_json=? WHERE id=?", (json.dumps(order), self.job_id))
        self.client.post(self.customer_url, data=self.customer_form, follow_redirects=False)
        response = self.client.get("/dashboard/automation")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Дата заявки:</b> 29.09.2026", response.text)
        self.assertIn("Заказчик:</b> ООО «Заказчик &lt;источник&gt;»", response.text)
        self.assertNotIn("Заказчик:</b> ООО «СТРОЙЛОГИСТИКА»", response.text)
        self.fetch.assert_not_called()

    def test_list_does_not_present_import_date_as_missing_order_date(self):
        response = self.client.get("/dashboard/automation")
        self.assertIn("Дата заявки:</b> не передана", response.text)
        self.assertIn("Заказчик:</b> не указан", response.text)

    def test_pending_and_manual_requests_show_metadata_before_supplier_search(self):
        pending = pipeline.save_api_order_for_manual_start({**self.order, "id": 999, "creationDate": "2026-10-01", "customer": {"name": "Заказчик без подбора"}})
        manual = pipeline.save_manual_order("Ручная заявка", [{"goodName": "Товар", "count": 1}], customer_name="Ручной заказчик")
        response = self.client.get("/dashboard/automation")
        self.assertIn("Дата заявки:</b> 01.10.2026", response.text)
        self.assertIn("Заказчик:</b> Заказчик без подбора", response.text)
        self.assertIn("Заказчик:</b> Ручной заказчик", response.text)
        self.assertNotEqual(pending["job_id"], manual["job_id"])

    def test_default_selection_distinguishes_suppliers_with_same_sku_and_allows_override(self):
        cheap = {"supplier": "КРЕП-КОМП", "sku": "shared", "name": "Товар", "price": 40, "match_score": 1}
        expensive = {**cheap, "supplier": "ВИ", "price": 100}
        result = copy.deepcopy(self.result)
        result["items"][0].update(selected=cheap, candidates=[cheap, expensive], selection_rule="lowest_unit_price")
        with pipeline._connect() as conn:
            conn.execute("UPDATE automation_jobs SET result_json=? WHERE id=?", (json.dumps(result), self.job_id))
        response = self.client.get(f"/dashboard/automation/jobs/{self.job_id}/review")
        self.assertIn("<option value='0' data-price='42.0' selected>", response.text)
        self.assertIn("<option value='1' data-price='105.0' >", response.text)
        response = self.client.post(f"/dashboard/automation/jobs/{self.job_id}/review", data={
            "candidate_1": "1", "quantity_1": "2", "price_1": "105", "include_1": "1",
        }, follow_redirects=False)
        self.assertEqual(response.status_code, 303)
        saved = pipeline.load_automation_offer_context(self.order["id"])["result"]["items"][0]
        self.assertEqual(saved["selected"]["supplier"], "ВИ")
        self.assertEqual(saved["selection_rule"], "operator_selected")


if __name__ == "__main__":
    unittest.main()
