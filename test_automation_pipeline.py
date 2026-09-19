import json
import os
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from io import BytesIO

import automation_pipeline as pipeline
from supplier_adapters import SupplierQuote


def email(order_id=37299999, message_id="mail-1"):
    return f"""From: zakupay@sel-be.ru
To: test@example.com
Message-ID: <{message_id}@example.com>
Subject: Запрос предложения по заявке {order_id}
Content-Type: text/plain; charset=utf-8

Новая заявка {order_id}
""".encode()


ORDER = {
    "id": 37299999,
    "name": "Тест",
    "delay": 0,
    "orderItems": [{"id": 10, "goodName": "Маркер черный 1 мм", "count": 10, "unit": {"name": "шт"}}],
}


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        pipeline.DB_PATH = os.path.join(self.tmp.name, "automation.db")

    def tearDown(self):
        self.tmp.cleanup()

    def test_search_variants_expand_russian_catalog_word_forms(self):
        variants = pipeline._search_variants("Нарукавники брезентовые")
        self.assertIn("нарукавники брезентовые", variants)
        self.assertIn("брезентовые нарукавники", variants)
        self.assertIn("нарукавник брезентовый", variants)
        self.assertIn("брезентовый нарукавник", variants)
        self.assertIn("нарукавники", variants)
        self.assertIn("36641496", variants)

    def test_purchase_label_explains_pack_and_unit_price_basis(self):
        packed = pipeline._purchase_label(
            {"name": "Саморезы, 1000 шт.", "price": 1339.0, "unit": "Упаковка"},
            "шт",
        )
        single = pipeline._purchase_label(
            {"name": "Защитные очки", "price": 76.0, "unit": "шт"},
            "шт",
        )
        self.assertIn("закупка ВИ: 1339 ₽ за упаковку 1000 шт.", packed)
        self.assertIn("закупка ВИ: 76 ₽ за 1 шт", single)

    def test_refresh_can_rebuild_order_from_saved_job(self):
        order = pipeline._order_from_saved_result(
            {"order_id": 37247138, "subject": "Заявка из письма"},
            {
                "order_name": "Брезентовые изделия",
                "items": [{
                    "order_item_id": 1,
                    "requested_name": "Нарукавники брезентовые",
                    "quantity": 25,
                    "unit": "пара",
                }],
            },
        )
        self.assertEqual(order["source"], "saved_automation_job")
        self.assertEqual(order["orderItems"][0]["goodName"], "Нарукавники брезентовые")
        self.assertEqual(order["orderItems"][0]["unit"]["name"], "пара")

    def test_refresh_prefers_full_saved_order_snapshot(self):
        row = {
            "id": 89,
            "order_id": 37247138,
            "subject": "Заявка из письма",
            "order_json": json.dumps({
                "id": 37247138,
                "deliveryAddress": "Сохранённый адрес",
                "orderItems": [{"goodName": "Нарукавники брезентовые", "count": 25}],
            }, ensure_ascii=False),
        }
        order = pipeline._saved_order_snapshot(row, {"items": []})
        self.assertEqual(order["source"], "saved_order_snapshot")
        self.assertEqual(order["deliveryAddress"], "Сохранённый адрес")

    @patch.object(pipeline.VseinstrumentiAdapter, "search")
    def test_broad_catalog_variant_recovers_product(self, search):
        order = {
            "id": 37299999,
            "orderItems": [{
                "id": 1,
                "goodName": "Нарукавники брезентовые",
                "count": 10,
                "unit": {"name": "пар"},
            }],
        }

        def results(query, limit=8):
            if query == "нарукавники":
                return [SupplierQuote(
                    supplier="ВИ",
                    name="Нарукавники брезентовые ФАЕР РЕЗИСТ п.420гр",
                    sku="36641496",
                    price=175,
                    stock=100,
                )]
            return []

        search.side_effect = results
        item = pipeline.build_vi_draft(order)["items"][0]
        self.assertEqual(item["selected"]["sku"], "36641496")
        self.assertEqual(item["decision"], "review")

    @patch.object(pipeline.VseinstrumentiAdapter, "search")
    def test_builds_draft_and_deduplicates_message(self, search):
        search.return_value = [SupplierQuote(
            supplier="ВсеИнструменты.ру", name="Маркер черный 1 мм",
            sku="123", price=100, stock=50,
        )]
        fetch = lambda order_id, force=False: ORDER if order_id == 37299999 else None
        first = pipeline.process_email(email(), fetch)
        second = pipeline.process_email(email(), fetch)
        self.assertFalse(first["duplicate"])
        self.assertTrue(second["duplicate"])
        self.assertEqual(first["result"]["summary"]["positions"], 1)
        self.assertEqual(first["result"]["items"][0]["proposed_unit_price"], 105.0)
        self.assertEqual(first["result"]["invoice_number"], 240)
        self.assertEqual(second["result"]["invoice_number"], 240)
        self.assertEqual(first["result"]["vat_rate"], 0.22)
        self.assertEqual(first["result"]["prepayment_percent"], 100.0)
        self.assertTrue(first["result"]["delivery_included"])
        self.assertFalse(first["result"]["live_offer_created"])
        self.assertTrue(first["result"]["zakupay_line_ids_complete"])
        with pipeline._connect() as conn:
            row = pipeline._execute(
                conn, "SELECT order_json FROM automation_jobs WHERE id=?", (first["job_id"],)
            ).fetchone()
        saved_order = json.loads(row["order_json"])
        self.assertEqual(saved_order["orderItems"][0]["goodName"], "Маркер черный 1 мм")

    @patch.object(pipeline.VseinstrumentiAdapter, "search")
    def test_assigns_sequential_invoice_numbers(self, search):
        search.return_value = [SupplierQuote(
            supplier="ВсеИнструменты.ру", name="Маркер черный 1 мм",
            sku="123", price=100, stock=50,
        )]
        fetch = lambda order_id, force=False: ORDER
        first = pipeline.process_email(email(message_id="sequence-1"), fetch)
        second = pipeline.process_email(email(message_id="sequence-2"), fetch)
        self.assertEqual(first["result"]["invoice_number"], 240)
        self.assertEqual(second["result"]["invoice_number"], 241)

    def test_missing_order_is_recorded_as_failed(self):
        with self.assertRaisesRegex(RuntimeError, "Заявка не получена"):
            pipeline.process_email(email(message_id="mail-2"), lambda *_args, **_kwargs: None)

    @patch.object(pipeline.VseinstrumentiAdapter, "search")
    def test_partial_invoice_contains_only_confirmed_rows(self, search):
        order = {
            "id": 37299999,
            "delay": 0,
            "orderItems": [
                {"id": 1, "goodName": "Маркер черный 1 мм", "count": 2, "unit": {"name": "шт"}},
                {"id": 2, "goodName": "Ковер 100x100 см", "count": 1, "unit": {"name": "шт"}},
            ],
        }

        def results(query, limit=8):
            if "Маркер" in query:
                return [SupplierQuote(supplier="ВИ", name="Маркер черный 1 мм", price=100, stock=10)]
            return [SupplierQuote(supplier="ВИ", name="Ковер 40x60 см", price=200, stock=10)]

        search.side_effect = results
        result = pipeline.process_email(email(message_id="partial"), lambda *_args, **_kwargs: order)
        self.assertEqual(result["status"], "ready_for_review")
        self.assertEqual(result["result"]["summary"]["included_in_invoice"], 1)
        self.assertEqual(result["result"]["summary"]["excluded_from_invoice"], 1)
        self.assertIn("attachment_base64", result["gmail_draft"])

        workbook = load_workbook(BytesIO(pipeline.build_invoice_xlsx(result["result"])))
        values = [cell.value for row in workbook.active.iter_rows() for cell in row]
        self.assertIn("Маркер черный 1 мм", values)
        self.assertNotIn("Ковер 40x60 см", values)

    @patch.object(pipeline.VseinstrumentiAdapter, "search")
    def test_pack_quantity_is_converted_before_stock_check(self, search):
        order = {
            "id": 37299999,
            "orderItems": [{
                "id": 1,
                "goodName": "Дюбель Tech-Krep 10x180 мм 130103",
                "count": 800,
                "unit": {"name": "шт"},
            }],
        }
        search.return_value = [SupplierQuote(
            supplier="ВИ",
            name="Дюбель Tech-Krep 10x180 мм, 200 шт 130103",
            article="130103",
            price=4000,
            stock=20,
            unit="Упаковка",
        )]
        draft = pipeline.build_vi_draft(order)
        item = draft["items"][0]
        self.assertEqual(item["pack_size"], 200)
        self.assertEqual(item["purchase_units"], 4)
        self.assertTrue(item["stock_confirmed"])
        self.assertEqual(item["decision"], "auto_ready")
        self.assertEqual(item["proposed_unit_price"], 21.0)

    @patch.object(pipeline.VseinstrumentiAdapter, "search")
    def test_delivery_date_allows_exact_item_without_numeric_stock(self, search):
        order = {
            "id": 37299999,
            "orderItems": [{"id": 1, "goodName": "Замок IEK YZK10-18-20-40", "count": 5, "unit": {"name": "шт"}}],
        }
        search.return_value = [SupplierQuote(
            supplier="ВИ",
            name="Замок IEK 18-20/40 YZK10-18-20-40",
            article="YZK10-18-20-40",
            price=1000,
            stock=None,
            courier_date="2026-09-18T03:00:00Z",
        )]
        item = pipeline.build_vi_draft(order)["items"][0]
        self.assertEqual(item["decision"], "auto_ready")
        self.assertEqual(item["availability_status"], "доступно к заказу, количество не подтверждено")

    @patch.object(pipeline.VseinstrumentiAdapter, "search")
    def test_conflicting_identifier_never_auto_accepted(self, search):
        order = {
            "id": 37299999,
            "orderItems": [{"id": 1, "goodName": "Контактор EKF km-2-32-20", "count": 1, "unit": {"name": "шт"}}],
        }
        search.return_value = [SupplierQuote(
            supplier="ВИ",
            name="Контактор EKF km-3-32-40",
            article="km-3-32-40",
            price=3000,
            stock=10,
        )]
        item = pipeline.build_vi_draft(order)["items"][0]
        self.assertNotEqual(item["decision"], "auto_ready")
        self.assertIn("не совпадает модель/артикул", item["replacement_details"])

    def test_operator_approved_row_is_included_in_partial_invoice(self):
        draft = {
            "invoice_number": 240,
            "order_id": 1,
            "items": [
                {"decision": "approved", "requested_name": "Товар 1", "selected": {"name": "Аналог 1"}, "quantity": 2, "unit": "шт", "proposed_unit_price": 100},
                {"decision": "excluded", "requested_name": "Товар 2", "selected": {"name": "Аналог 2"}, "quantity": 1, "unit": "шт", "proposed_unit_price": 200},
            ],
        }
        workbook = load_workbook(BytesIO(pipeline.build_invoice_xlsx(draft)))
        values = [cell.value for row in workbook.active.iter_rows() for cell in row]
        self.assertIn("Аналог 1", values)
        self.assertNotIn("Аналог 2", values)

    @patch.object(pipeline.VseinstrumentiAdapter, "search")
    def test_api_order_is_persisted_and_deduplicated_by_order_id(self, search):
        search.return_value = [SupplierQuote(
            supplier="ВИ", name="Маркер черный 1 мм", sku="123", price=100, stock=50,
        )]
        first = pipeline.process_api_order(ORDER)
        second = pipeline.process_api_order(ORDER)
        self.assertFalse(first["duplicate"])
        self.assertTrue(second["duplicate"])
        self.assertEqual(first["result"]["invoice_number"], 240)
        context = pipeline.load_automation_offer_context(ORDER["id"])
        self.assertEqual(context["job_id"], first["job_id"])
        self.assertEqual(context["order"]["orderItems"][0]["id"], 10)

        pipeline.mark_automation_offer_created(first["job_id"], offer_id="offer-1", file_id="file-1")
        context = pipeline.load_automation_offer_context(ORDER["id"])
        self.assertTrue(context["result"]["live_offer_created"])
        self.assertEqual(context["result"]["live_offer_id"], "offer-1")

    @patch.object(pipeline.VseinstrumentiAdapter, "search")
    def test_missing_line_ids_are_enriched_once_and_persisted(self, search):
        search.return_value = [SupplierQuote(
            supplier="ВИ", name="Маркер черный 1 мм", sku="123", price=100, stock=50,
        )]
        order_without_ids = {
            "id": ORDER["id"], "name": "Тест", "delay": 0,
            "orderItems": [{"goodName": "Маркер черный 1 мм", "count": 10, "unit": {"name": "шт"}}],
        }
        created = pipeline.process_api_order(order_without_ids)
        context = pipeline.enrich_automation_offer_context(ORDER["id"], ORDER)
        self.assertEqual(context["job_id"], created["job_id"])
        self.assertEqual(context["result"]["items"][0]["order_item_id"], 10)
        self.assertTrue(context["result"]["order_id_enrichment_found"])
        self.assertEqual(context["order"]["orderItems"][0]["id"], 10)

    @patch.object(pipeline, "build_vi_draft")
    def test_failed_api_order_can_be_retried(self, build):
        build.side_effect = RuntimeError("temporary")
        with self.assertRaisesRegex(RuntimeError, "temporary"):
            pipeline.process_api_order(ORDER)
        build.side_effect = None
        build.return_value = {
            "status": "ready_for_review",
            "summary": {"auto_ready": 1},
            "invoice_number": None,
        }
        retried = pipeline.process_api_order(ORDER)
        self.assertFalse(retried["duplicate"])
        self.assertEqual(retried["status"], "ready_for_review")


if __name__ == "__main__":
    unittest.main()
