"""Porcelain catalogue snapshots, explainable analogues and delivery estimates."""
from __future__ import annotations

import hashlib
import io
import json
import math
import re
import zipfile
from collections import defaultdict
from datetime import datetime, timezone
from decimal import Decimal, ROUND_CEILING, ROUND_HALF_UP
from functools import lru_cache

from invoice_generator import SELLER, customer_validation_error, validate_seller
from price_lists import database, execute, norm, number

MAX_IMPORT = 12 * 1024 * 1024


def schema_statements(postgres):
    serial, blob = ("BIGSERIAL", "BYTEA") if postgres else ("INTEGER", "BLOB")
    auto = "" if postgres else " AUTOINCREMENT"
    return [
        f"""CREATE TABLE IF NOT EXISTS tile_catalogs (
            id {serial} PRIMARY KEY{auto}, content_hash TEXT NOT NULL UNIQUE,
            catalog_json TEXT NOT NULL, archive_bytes {blob} NOT NULL, imported_at TEXT NOT NULL)""",
        f"""CREATE TABLE IF NOT EXISTS tile_quotes (
            id {serial} PRIMARY KEY{auto}, job_id BIGINT NOT NULL, catalog_id BIGINT NOT NULL,
            snapshot_json TEXT NOT NULL, created_at TEXT NOT NULL)""",
        "CREATE INDEX IF NOT EXISTS idx_tile_quotes_job ON tile_quotes(job_id,id)",
        """CREATE TABLE IF NOT EXISTS tile_sellers (
            id TEXT PRIMARY KEY, profile_json TEXT NOT NULL, updated_at TEXT NOT NULL)""",
    ]


def initialize_schema(conn, postgres):
    for sql in schema_statements(postgres):
        conn.execute(sql)


def now():
    return datetime.now(timezone.utc).isoformat()


def _json(value):
    return json.dumps(value, ensure_ascii=False, allow_nan=False)


def _positive(value, label, zero=False):
    parsed = number(value)
    if parsed is None or (parsed < 0 if zero else parsed <= 0):
        raise ValueError(f"{label}: введите {'неотрицательное' if zero else 'положительное'} число")
    return parsed


def money(value):
    return float(Decimal(str(value)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP))


def validate_archive(content):
    if not content or len(content) > MAX_IMPORT:
        raise ValueError("Архив должен быть ZIP размером до 12 МБ")
    try:
        z = zipfile.ZipFile(io.BytesIO(content))
        files = z.infolist()
        if len(files) > 2000 or sum(f.file_size for f in files) > 60 * 1024 * 1024:
            raise ValueError("Слишком большой распакованный архив")
        if len({f.filename for f in files}) != len(files):
            raise ValueError("В архиве повторяются имена файлов")
        if z.getinfo('catalog.json').file_size > 4 * 1024 * 1024:
            raise ValueError("Слишком большой каталог")
        raw = z.read('catalog.json')
        data = json.loads(raw)
        products = data.get('products')
        if not isinstance(products, list) or not products or len(products) > 10000:
            raise ValueError("В catalog.json нет списка товаров")
        ids = set()
        for p in products:
            if not isinstance(p, dict):
                raise ValueError('Каждая позиция каталога должна быть объектом')
            key = str(p.get('id') or '')
            if not re.fullmatch(r'[A-Za-z0-9_-]{1,80}', key) or key in ids or not p.get('name'):
                raise ValueError("Нужны уникальные ID и названия товаров")
            ids.add(key)
            image = p.get('image')
            if image and (not re.fullmatch(r'assets/[A-Za-z0-9_.-]+\.(jpg|jpeg|png)', image, re.I) or image not in z.namelist()):
                raise ValueError("Не найдено изображение товара в assets")
        return data
    except (zipfile.BadZipFile, KeyError, TypeError, json.JSONDecodeError) as exc:
        raise ValueError("Нужен архив с catalog.json и изображениями assets") from exc


def import_catalog(content):
    data = validate_archive(content)
    digest = hashlib.sha256(content).hexdigest()
    with database() as conn:
        execute(conn, """INSERT INTO tile_catalogs(content_hash,catalog_json,archive_bytes,imported_at)
            VALUES (?,?,?,?) ON CONFLICT(content_hash) DO NOTHING""", (digest, _json(data), content, now()))
        row = execute(conn, "SELECT id FROM tile_catalogs WHERE content_hash=?", (digest,)).fetchone()
    return int(row['id'])


def import_workbook(content):
    """Overlay the user's explicit catalogue columns, retaining image identities."""
    from openpyxl import load_workbook
    previous = load_catalog(with_archive=True)
    if not previous:
        raise ValueError('Сначала загрузите ZIP с каталогом и фотографиями')
    if len(content) > MAX_IMPORT:
        raise ValueError('Размер таблицы должен быть до 12 МБ')
    data = json.loads(previous['catalog_json'])
    products = {p['id']: p for p in data['products']}
    warnings = []
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(f.file_size for f in archive.infolist()) > 60*1024*1024:
                raise ValueError('Слишком большая таблица')
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = book['Каталог']
        if sheet.cell(5, 2).value != 'ID каталога' or sheet.cell(5, 6).value != 'Цена, ₽/м²':
            raise ValueError('Нужен лист «Каталог» рабочей таблицы КерамаХолл с неизменёнными заголовками')
        seen = set()
        columns = ['name','format','stock_m2','price_rub_m2','vat','color','pattern','thickness_mm','surface',
                   'pieces_box','box_m2','pallet_m2','kg_m2_estimate','kg_m2_confirmed','pallet_m2_estimate']
        numeric = {'stock_m2','price_rub_m2','thickness_mm','pieces_box','box_m2','pallet_m2','kg_m2_estimate','kg_m2_confirmed','pallet_m2_estimate'}
        # Google exports can omit worksheet dimensions in read-only mode.
        for values in sheet.iter_rows(min_row=6, max_row=min(sheet.max_row or 10005, 10005), max_col=20, values_only=True):
            key = str(values[1] or '').strip()
            if not key: continue
            if key not in products or key in seen:
                raise ValueError('Неизвестный или повторный ID в таблице: ' + key)
            seen.add(key)
            p = products[key]
            for field, value in zip(columns, values[2:17]):
                if field == 'name' and str(value or '').startswith(('#gid=', 'http://', 'https://')):
                    warnings.append(key + ': вместо названия в таблице ссылка; сохранено название по ID из каталога')
                    continue
                if field in numeric:
                    parsed = number(value)
                    if value not in (None, '') and (parsed is None or parsed < 0):
                        raise ValueError(f'{key}: некорректное значение {field}')
                    p[field] = parsed
                else:
                    p[field] = str(value or '').strip() or None
        book.close()
        if not seen: raise ValueError('Нет строк каталога')
    except (zipfile.BadZipFile, KeyError) as exc:
        raise ValueError('Нужен XLSX с листом «Каталог»') from exc
    data['workbook_imported_at'] = now()
    data['workbook_import_warnings'] = warnings
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(bytes(previous['archive_bytes']))) as source, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            target.writestr(item.filename, _json(data).encode() if item.filename == 'catalog.json' else source.read(item.filename))
    return import_catalog(output.getvalue())


def load_catalog(catalog_id=None, with_archive=False):
    fields = 'id,catalog_json,imported_at' + (',archive_bytes' if with_archive else '')
    with database() as conn:
        row = execute(conn, f"SELECT {fields} FROM tile_catalogs" + (' WHERE id=?' if catalog_id else ' ORDER BY id DESC LIMIT 1'),
                      (catalog_id,) if catalog_id else ()).fetchone()
    if not row:
        return None
    return {**dict(row), 'data': json.loads(row['catalog_json'])}


@lru_cache(maxsize=2)
def _catalog_assets(database_identity, catalog_id):
    # Immutable versions: load the archive once, not once per product image.
    catalog = load_catalog(catalog_id, with_archive=True)
    if not catalog:
        raise LookupError('Каталог не найден')
    with zipfile.ZipFile(io.BytesIO(bytes(catalog['archive_bytes']))) as z:
        return {p['id']: z.read(p['image']) for p in catalog['data']['products'] if p.get('image')}


def catalogue_image(catalog_id, product_id):
    from automation_pipeline import _database_key
    try:
        return _catalog_assets(_database_key(), catalog_id).get(product_id)
    except LookupError:
        return None


def is_tile(text):
    text = norm(text)
    if re.search(r'\b(клей|затирк\w*|крестик\w*|сверл\w*|сверел|диск\w*|плиткорез\w*|система\s+выравнив\w*)\b', text):
        return False
    return bool(re.search(r'керамогранит|porcelain|керамическ\w*\s+плитк|плитк\w*\s+керамическ', text))


GROUPS = {
    'color': {'белый': r'\bбел(?:ый|ая|ое|ые|ого|ой|ую|ых|ыми|ом|ому)\b|white|вайт', 'серый': r'\bсер(?:ый|ая|ое|ые|ого|ой|ую|ых|ыми|ом|ому)\b|grey|gray|грей', 'бежевый': r'беж|beige',
              'черный': r'черн|black|блэк', 'коричневый': r'коричн|brown', 'кремовый': r'крем|cream',
              'антрацит': r'антрацит|anthracite'},
    'pattern': {'мрамор': r'мрамор|marble|калакт|calacat', 'бетон': r'бетон|concrete',
                'дерево': r'дерев|дуб|wood', 'камень': r'камень|камен|stone',
                'травертин': r'траверт|travert', 'орнамент': r'орнамент|пэчворк|patchwork',
                'однотонный': r'однотон|моноколор|monocolor'},
    'surface': {'матовая': r'матов|матт|matt|неполирован', 'полированная': r'(?<!не)полирован|polished|глянц',
                'лаппатированная': r'лаппат|lappat'},
}


def requirements(text):
    text = str(text or '').casefold().replace('ё', 'е')
    result = {key: {name for name, pattern in groups.items() if re.search(pattern, text)} for key, groups in GROUPS.items()}
    size = re.search(r'(\d+(?:[.,]\d+)?)\s*[xх×*]\s*(\d+(?:[.,]\d+)?)(?:\s*[xх×*]\s*(\d+(?:[.,]\d+)?))?\s*(мм|см)?', text)
    result['format'] = None
    result['thickness'] = None
    if size:
        values = [float(size.group(i).replace(',', '.')) for i in (1, 2)]
        divisor = 10 if (size.group(4) == 'мм' or (size.group(4) != 'см' and max(values) >= 200)) else 1
        result['format'] = tuple(sorted(round(n / divisor, 3) for n in values))
        if size.group(3):
            result['thickness'] = float(size.group(3).replace(',', '.'))
        text = text[:size.start()] + text[size.end():]
    thick = re.search(r'(?:толщ\w*\s*)?(\d+(?:[.,]\d+)?)\s*мм', text)
    if thick:
        result['thickness'] = float(thick.group(1).replace(',', '.'))
    result['tone'] = {'темный' for _ in [1] if re.search(r'темн|dark', text)} | {'светлый' for _ in [1] if re.search(r'светл|light', text)}
    return result


def product_requirements(p):
    result = requirements(' '.join(str(p.get(k) or '') for k in ('name', 'color', 'pattern', 'surface')))
    result['format'] = requirements(p.get('format'))['format']
    result['thickness'] = number(p.get('thickness_mm'))
    return result


def assess(requested, product):
    wanted, actual = requirements(requested), product_requirements(product)
    conflicts, unknown, reasons = [], [], []
    score = 0
    for key, label, weight in [('format', 'Формат', 30), ('thickness', 'Толщина', 10),
                               ('color', 'Цвет', 25), ('pattern', 'Рисунок', 25), ('surface', 'Поверхность', 10), ('tone', 'Тон', 10)]:
        w, a = wanted[key], actual[key]
        if not w:
            continue
        if not a:
            unknown.append(label + ' не указан в каталоге')
        elif (w != a if key in {'format', 'thickness'} else not w.intersection(a)):
            conflicts.append(label + ' отличается')
        else:
            score += weight
            reasons.append(label + ' совпадает')
    if norm(product.get('name')) and norm(product['name']) in norm(requested):
        score += 40
        reasons.append('Название найдено в заявке')
    return {'score': score, 'reasons': reasons, 'unknown': unknown, 'conflicts': conflicts}


def find_analogues(data, requested, limit=8):
    found = []
    for p in data['products']:
        check = assess(requested, p)
        if check['conflicts']:
            continue
        found.append({'product': p, **check})
    return sorted(found, key=lambda r: (-r['score'], r['product'].get('stock_m2') is None,
                                      not bool(r['product'].get('image')), r['product']['name']))[:limit]


def delivery_address(order):
    value = order.get('deliveryAddress') or order.get('address') or ''
    if isinstance(value, dict):
        return str(value.get('fullAddress') or value.get('address') or value.get('name') or '')
    return str(value)


def sellers():
    with database() as conn:
        rows = execute(conn, 'SELECT id,profile_json FROM tile_sellers ORDER BY id').fetchall()
    return {'avior': dict(SELLER), **{r['id']: json.loads(r['profile_json']) for r in rows}}


def save_seller(values):
    profile = validate_seller(values)
    key = 'firm-' + profile['inn'] + '-' + profile['account'][-6:]
    with database() as conn:
        execute(conn, """INSERT INTO tile_sellers(id,profile_json,updated_at) VALUES (?,?,?)
            ON CONFLICT(id) DO UPDATE SET profile_json=excluded.profile_json,updated_at=excluded.updated_at""", (key, _json(profile), now()))
    return key


def make_quote(order, catalog, form):
    products = {p['id']: p for p in catalog['data']['products']}
    lines, issues, stock_used = [], [], defaultdict(float)
    for index, item in enumerate(order.get('orderItems') or [], 1):
        key = str(form.get(f'product_{index}') or '')
        if not key:
            continue
        if key not in products:
            raise ValueError(f'Позиция {index}: товар не найден')
        p = products[key]
        requested = str(item.get('goodName') or item.get('name') or '')
        check = assess(requested + ' ' + str(form.get(f'features_{index}') or ''), p)
        if check['conflicts']:
            raise ValueError(f"Позиция {index}: {', '.join(check['conflicts'])}")
        area = _positive(form.get(f'area_{index}'), f'Позиция {index}: площадь, м²')
        reserve = _positive(form.get(f'reserve_{index}') or '0', 'Запас, %', zero=True)
        box = number(form.get(f'box_{index}')) if form.get(f'box_{index}') not in (None, '') else number(p.get('box_m2'))
        if box is not None and box <= 0:
            raise ValueError(f'Позиция {index}: площадь коробки должна быть положительной')
        boxes = math.ceil(Decimal(str(area)) * (1 + Decimal(str(reserve))/100) / Decimal(str(box))) if box and box > 0 else None
        shipped = float(Decimal(str(boxes)) * Decimal(str(box))) if boxes else None
        price = number(form.get(f'price_{index}'))
        if price is not None and price <= 0:
            raise ValueError(f'Позиция {index}: цена должна быть больше нуля')
        pallet_area = number(form.get(f'pallet_{index}')) if form.get(f'pallet_{index}') not in (None, '') else number(p.get('pallet_m2')) or number(p.get('pallet_m2_estimate'))
        kg = number(form.get(f'kg_{index}')) if form.get(f'kg_{index}') not in (None, '') else number(p.get('kg_m2_confirmed')) or number(p.get('kg_m2_estimate'))
        if (kg is not None and kg <= 0) or (pallet_area is not None and pallet_area <= 0):
            raise ValueError(f'Позиция {index}: вес и площадь паллеты должны быть положительными')
        pallets = math.ceil(Decimal(str(shipped)) / Decimal(str(pallet_area))) if shipped and pallet_area and pallet_area > 0 else None
        weight = shipped * kg + pallets * (number(p.get('pallet_kg_estimate')) or 30) if shipped and kg and pallets else None
        missing = []
        if not shipped: missing.append('Нет площади коробки')
        if price is None: missing.append('Нужна продажная цена за м²')
        if p.get('stock_m2') is None: missing.append('Наличие неизвестно')
        if form.get(f'approve_{index}') != 'yes': missing.append('Подтвердите аналог и наличие')
        if not weight: missing.append('Недостаточно данных о весе/паллетах')
        if shipped: stock_used[key] += shipped
        lines.append({'position': index, 'requested_name': requested, 'product': dict(p), 'required_m2': area,
                      'reserve_pct': reserve, 'boxes': boxes, 'quantity': shipped, 'unit': 'м²', 'pallets': pallets,
                      'weight_kg': weight, 'box_m2_used': box, 'kg_m2_used': kg, 'pallet_m2_used': pallet_area,
                      'proposed_unit_price': price, 'amount': money(Decimal(str(shipped))*Decimal(str(price))) if shipped and price else None,
                      'approved': form.get(f'approve_{index}') == 'yes', 'check': check, 'issues': missing})
    if not lines:
        raise ValueError('Выберите хотя бы один товар для предложения')
    for line in lines:
        stock = number(line['product'].get('stock_m2'))
        if stock is not None and stock_used[line['product']['id']] > stock + 0.000001:
            line['issues'].append('Недостаточный остаток по сумме всех строк')
        issues.extend(f"Позиция {line['position']}: {error}" for error in line['issues'])
    address = str(form.get('address') or '').strip()
    if not address: issues.append('Укажите адрес доставки')
    zone = str(form.get('zone') or '')
    km = number(form.get('km'))
    unload = number(form.get('unloading'))
    if zone not in {'mkad', 'ttk', 'outside'}: issues.append('Подтвердите зону доставки: только Москва и МО')
    if zone == 'outside' and (km is None or km < 0): issues.append('Укажите километры за МКАД по маршруту')
    if unload is None or unload < 0: issues.append('Укажите согласованную разгрузку; 0 — разгружает клиент')
    weight = sum(l['weight_kg'] or 0 for l in lines) if all(l['weight_kg'] for l in lines) else None
    shipping = None
    tariffs = catalog['data'].get('shipping', {}).get('tariffs', [])
    vehicle = str(form.get('vehicle') or '')
    if weight and zone in {'mkad', 'ttk', 'outside'} and unload is not None and unload >= 0 and (zone != 'outside' or (km is not None and km >= 0)):
        options = []
        for t in tariffs:
            if vehicle and t['vehicle'] != vehicle: continue
            if not vehicle and 'манипулятор' in norm(t['vehicle']): continue
            trips = math.ceil(Decimal(str(weight)) / Decimal(str(_positive(t.get('capacity_kg'), 'Грузоподъёмность'))))
            costs = {band: money(trips * (float(t[band]) + (km if zone == 'outside' else 0) * float(t['km_' + band]) + (float(t.get('center') or 0) if zone == 'ttk' else 0)) + unload) for band in ('low', 'mid', 'high')}
            options.append({'vehicle': t['vehicle'], 'trips': trips, **costs})
        if options:
            shipping = min(options, key=lambda option: option['mid'])
    if shipping is None: issues.append('Доставка не рассчитана')
    sale_delivery = number(form.get('sale_delivery'))
    if sale_delivery is None and shipping:
        sale_delivery = float((Decimal(str(shipping['mid']))/500).to_integral_value(rounding=ROUND_CEILING)*500)
    if sale_delivery is not None and sale_delivery < 0:
        raise ValueError('Продажная стоимость доставки должна быть неотрицательной')
    if form.get('transport_confirmed') != 'yes': issues.append('Подтвердите перевозку, габариты и окончательную стоимость доставки')
    vat = str(form.get('vat') or '')
    if vat not in {'0','10','20','22'}: issues.append('Выберите НДС нашей фирмы')
    seller = sellers().get(str(form.get('seller') or ''))
    if not seller: issues.append('Выберите нашу фирму')
    if form.get('seller_confirmed') != 'yes': issues.append('Подтвердите реквизиты нашей фирмы')
    customer = dict(order.get('customer') or {})
    for field, key in [('customer_name','name'),('customer_inn','inn'),('customer_kpp','kpp'),('customer_address','legalAddress')]:
        if field in form:
            customer[key] = str(form.get(field) or '').strip()
            if key == 'name':
                customer.pop('fullName', None)
                customer.pop('shortName', None)
    error = customer_validation_error(customer)
    if error: issues.append(error)
    if not str(form.get('invoice_number') or '').strip(): issues.append('Укажите номер счёта')
    total = money(sum(l['amount'] for l in lines) + sale_delivery) if all(l['amount'] is not None for l in lines) and sale_delivery is not None else None
    return {'order_id': order.get('id'), 'customer': customer, 'seller': seller, 'vat_rate': int(vat)/100 if vat.isdigit() else None,
            'invoice_number': str(form.get('invoice_number') or '').strip(), 'address': address, 'lines': lines,
            'weight_kg': weight, 'shipping': shipping, 'delivery_price': sale_delivery, 'total': total,
            'issues': list(dict.fromkeys(issues)), 'inputs': dict(form), 'stock_as_of': catalog['data'].get('stock_as_of'),
            'catalog_id': catalog['id'], 'order_snapshot': order, 'created_at': now()}


def save_quote(job_id, catalog, snapshot):
    with database() as conn:
        row = execute(conn, 'INSERT INTO tile_quotes(job_id,catalog_id,snapshot_json,created_at) VALUES (?,?,?,?) RETURNING id',
                      (job_id, catalog['id'], _json(snapshot), now())).fetchone()
    return int(row['id'])


def load_quote(quote_id):
    with database() as conn:
        row = execute(conn, 'SELECT * FROM tile_quotes WHERE id=?', (quote_id,)).fetchone()
    return {**dict(row), 'snapshot': json.loads(row['snapshot_json'])} if row else None


def invoice_draft(snapshot):
    if snapshot['issues']:
        raise ValueError('Счёт пока недоступен: ' + '; '.join(snapshot['issues']))
    rows = [{'decision': 'approved', 'requested_name': line['requested_name'], 'selected': {'name': line['product']['name'] + ' ' + str(line['product'].get('format') or '')},
             'quantity': line['quantity'], 'unit': 'м²', 'proposed_unit_price': line['proposed_unit_price']} for line in snapshot['lines']]
    rows.append({'decision': 'approved', 'requested_name': 'Доставка по адресу: ' + snapshot['address'], 'quantity': 1,
                 'unit': 'услуга', 'proposed_unit_price': snapshot['delivery_price']})
    return {**snapshot, 'items': rows, 'payment_note': 'Оплата: 100% предоплата. Доставка указана отдельной строкой.'}
