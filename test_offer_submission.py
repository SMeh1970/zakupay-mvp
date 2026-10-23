"""Invoice deadlines must release the website without retrying a live offer."""

import asyncio
import copy
import html
import threading
import unittest
from unittest.mock import Mock, patch

import httpx
import requests
from fastapi import FastAPI
from fastapi.testclient import TestClient

from offer_panel_safe import install_offer_panel
from invoice_generator import SELLER, build_invoice_xlsx
from io import BytesIO
from openpyxl import load_workbook


class OfferSubmissionTests(unittest.TestCase):
    def setUp(self):
        self.sv = {**SELLER, 'name': 'ООО «СВ»', 'inn': '5638073633', 'kpp': '561101001', 'account': '4' * 20}
        self.accounts = [{'id': 'bank', 'accountNumber': SELLER['account'], 'company': {'inn': SELLER['inn'], 'name': SELLER['name']}},
                         {'id': 'sv-bank', 'accountNumber': self.sv['account'], 'company': {'inn': self.sv['inn'], 'name': self.sv['name']}}]
        get = patch('offer_panel_safe.requests.get', side_effect=lambda *a, **kw: self.response({'accounts': self.accounts}))
        get.start()
        self.addCleanup(get.stop)
        self.save_settings = Mock()
        self.order = {"id": 37463885, "orderItems": [{"id": 10, "goodName": "Товар", "count": 2}]}
        self.context = {
            "job_id": 456, "order": self.order,
            "result": {
                "customer": {"name": "ООО Покупатель", "inn": "7716997861", "kpp": "770801001"},
                "items": [{"order_item_id": 10, "quantity": 2, "unit": "шт", "decision": "approved"}],
            },
        }
        self.fetch = Mock(return_value=copy.deepcopy(self.order))
        self.claim = Mock(return_value={"claimed": True, "id": 1, "status": "sending"})
        self.submission = None
        self.finish = Mock(side_effect=self._finish)
        self.invoice = Mock(return_value=b"invoice")
        self.form = {
            "confirm_send": "SEND", "producer_offer_number": "278", "producer_offer_date": "2026-10-22",
            "destination_account_id": "bank", "use_10": "1", "qty_10": "2", "price_10": "100",
            "unit_10": "шт", "provider_name_10": "Товар", "vat_rate": "0.22",
        }
        self.app = FastAPI()

        @self.app.get("/probe")
        async def probe():
            return {"ok": True}

        install_offer_panel(
            self.app, Mock(), lambda: {}, "https://example.invalid", lambda v: html.escape(str(v)),
            fetch_order_by_id=self.fetch, load_offer_context=lambda oid: copy.deepcopy(self.context),
            begin_submission=self.claim, finish_submission=self.finish,
            latest_submission=lambda job_id: self.submission,
            order_hash=lambda order: str(order["orderItems"]), build_invoice=self.invoice,
            load_sellers=lambda: {'avior': SELLER, 'sv': self.sv}, save_invoice_settings=self.save_settings,
        )
        self.url = "/dashboard/order/37463885/offer/submit"
        self.client = TestClient(self.app)
        self.addCleanup(self.client.close)

    def _finish(self, submission_id, status, **kwargs):
        self.submission = {"id": submission_id, "status": status, **kwargs}

    @staticmethod
    def response(data, status=200):
        return Mock(status_code=status, ok=status < 400, headers={"content-type": "application/json"}, json=lambda: data)

    def upload(self):
        return self.response({"file": {"id": "file-123"}})

    def check_deadline(self, stage, post):
        """The blocked network thread cannot resume the rest of the submission."""
        started, release, completed = threading.Event(), threading.Event(), threading.Event()

        def slow_reply(*args, **kwargs):
            started.set()
            release.wait(2)
            completed.set()
            return copy.deepcopy(self.order) if stage == "order_check" else self.response({"id": "late-id"})

        if stage == "order_check":
            self.fetch.side_effect = slow_reply
            deadline_name = "ORDER_CHECK_DEADLINE_SECONDS"
        elif stage == "file_upload":
            post.side_effect = slow_reply
            deadline_name = "UPLOAD_DEADLINE_SECONDS"
        else:
            post.side_effect = [self.upload(), Mock()]

            def dispatch(*args, **kwargs):
                return self.upload() if "files" in kwargs else slow_reply(*args, **kwargs)

            post.side_effect = dispatch
            deadline_name = "OFFER_POST_DEADLINE_SECONDS"

        async def scenario():
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=self.app), base_url="http://test") as client:
                pending = asyncio.create_task(client.post(self.url, data=self.form))
                try:
                    self.assertTrue(await asyncio.to_thread(started.wait, 1))
                    # This must work while the upstream request is still blocked.
                    probe = await asyncio.wait_for(client.get("/probe"), timeout=0.3)
                    self.assertEqual(probe.status_code, 200)
                    self.assertFalse(release.is_set())
                    response = await asyncio.wait_for(pending, timeout=0.5)
                    if stage == "offer_post":
                        self.assertEqual(response.status_code, 409)
                        self.assertEqual(self.submission["status"], "unknown")
                        again = await client.post(self.url, data=self.form)
                        self.assertEqual(again.status_code, 409)
                        self.assertEqual(post.call_count, 2)  # One upload and exactly one offer POST.
                    else:
                        self.assertEqual(response.status_code, 503)
                        self.assertIn("Счёт не отправлен", response.text)
                        self.claim.assert_not_called()
                finally:
                    release.set()
                    await asyncio.to_thread(completed.wait, 1)
                    if not pending.done():
                        await pending

        try:
            with patch("offer_panel_safe." + deadline_name, 0.05):
                asyncio.run(scenario())
        finally:
            release.set()
        self.assertEqual(post.call_count, {"order_check": 0, "file_upload": 1, "offer_post": 2}[stage])

    @patch("offer_panel_safe.requests.post")
    def test_slow_preflight_leaves_site_responsive_and_late_reply_never_sends(self, post):
        self.check_deadline("order_check", post)
        self.invoice.assert_not_called()

    @patch("offer_panel_safe.requests.post")
    def test_slow_upload_never_continues_to_offer_post(self, post):
        self.check_deadline("file_upload", post)

    @patch("offer_panel_safe.requests.post")
    def test_slow_offer_post_blocks_retries_even_after_late_reply(self, post):
        self.check_deadline("offer_post", post)

    @patch("offer_panel_safe.requests.post")
    def test_success_confirms_offer_only_after_id_is_received(self, post):
        post.side_effect = [self.upload(), self.response({"id": "offer-123"})]
        response = self.client.post(self.url, data=self.form)
        self.assertEqual(response.status_code, 200)
        self.assertIn("Предложение создано", response.text)
        self.assertEqual(self.submission["status"], "confirmed")
        self.assertEqual(self.submission["offer_id"], "offer-123")
        self.assertEqual(post.call_count, 2)

    @patch("offer_panel_safe.requests.post")
    def test_actual_offer_id_response_is_confirmed_in_json_and_xml(self, post):
        replies = [self.response({'result': True, 'offerId': 340932504,
                                   'offerUrl': 'https://zakupay.sel-be.ru/core/offers/340932504/report'}),
                   self.response({'result': True, 'id': 'unrelated-id', 'offerId': 340932504}),
                   Mock(status_code=200, ok=True, headers={'content-type': 'application/xml'},
                        text='<response><result>true</result><offerId>340932504</offerId></response>')]
        for reply in replies:
            with self.subTest(reply=reply):
                self.submission = None
                post.side_effect = [self.upload(), reply]
                response = self.client.post(self.url, data=self.form)
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual(self.submission['status'], 'confirmed')
                self.assertEqual(self.submission['offer_id'], '340932504')

    @patch("offer_panel_safe.requests.post")
    def test_foreign_file_id_and_negative_result_do_not_confirm_an_offer(self, post):
        for data in ({'result': True, 'files': [{'id': 'only-file-id'}]},
                     {'result': False, 'offerId': 340932504}, {'result': True, 'offerId': 0},
                     {'result': True, 'offerId': 0, 'id': 'unrelated-id'}):
            with self.subTest(data=data):
                self.submission = None
                post.side_effect = [self.upload(), self.response(data)]
                self.assertEqual(self.client.post(self.url, data=self.form).status_code, 409)
                self.assertEqual(self.submission['status'], 'unknown')

    @patch("offer_panel_safe.requests.post")
    def test_network_failure_during_post_is_unknown_and_never_retried(self, post):
        post.side_effect = [self.upload(), requests.Timeout("slow")]
        self.assertEqual(self.client.post(self.url, data=self.form).status_code, 409)
        self.assertEqual(self.submission["status"], "unknown")
        self.assertEqual(self.client.post(self.url, data=self.form).status_code, 409)
        self.assertEqual(post.call_count, 2)

    @patch("offer_panel_safe.requests.post")
    def test_upstream_server_error_does_not_permit_duplicate_offer(self, post):
        post.side_effect = [self.upload(), self.response({"error": "upstream"}, 502)]
        self.assertEqual(self.client.post(self.url, data=self.form).status_code, 409)
        self.assertEqual(self.submission["status"], "unknown")

    @patch("offer_panel_safe.requests.post")
    def test_success_without_offer_id_is_not_reported_as_confirmed(self, post):
        for data in ({"result": True}, {"result": False, "id": "not-an-offer"}):
            with self.subTest(data=data):
                self.submission = None
                post.side_effect = [self.upload(), self.response(data)]
                response = self.client.post(self.url, data=self.form)
                self.assertEqual(response.status_code, 409)
                self.assertEqual(self.submission["status"], "unknown")

    @patch("offer_panel_safe.requests.post")
    def test_invalid_form_is_rejected_before_upload(self, post):
        response = self.client.post(self.url, data={**self.form, "destination_account_id": ""})
        self.assertEqual(response.status_code, 400)
        post.assert_not_called()

    @patch('offer_panel_safe.requests.post')
    def test_selected_sv_is_in_uploaded_workbook_and_saved_with_vat(self, post):
        self.context['result']['items'][0].update(requested_name='Товар', selected={'name': 'Товар'}, proposed_unit_price=100)
        self.invoice.side_effect = build_invoice_xlsx
        post.side_effect = [self.upload(), self.response({'id': 'offer-sv'})]
        response = self.client.post(self.url, data={**self.form, 'seller_id': 'sv', 'destination_account_id': 'sv-bank', 'vat_rate': '0.2'})
        self.assertEqual(response.status_code, 200, response.text)
        upload = post.call_args_list[0].kwargs['files']['file']
        workbook = load_workbook(BytesIO(upload[1]))
        text = ' '.join(str(cell.value or '') for row in workbook.active for cell in row)
        self.assertIn(self.sv['name'], text)
        self.assertIn(self.sv['inn'], text)
        self.assertNotIn(SELLER['inn'], text)
        self.assertNotIn('АВИОР', text)
        self.assertNotIn('AVIOR', upload[0])
        self.save_settings.assert_called_once_with(456, self.sv, 'sv-bank', 0.2)
        self.assertEqual(self.invoice.call_args.args[0]['vat_rate'], 0.2)

    @patch('offer_panel_safe.requests.post')
    def test_attachment_matches_selected_rows_edits_date_and_submitted_totals(self, post):
        self.context['result']['items'][0].update(selected={'name': 'Старое название'}, proposed_unit_price=100)
        self.order['orderItems'].append({'id': 11, 'goodName': 'Исключённый товар', 'count': 5})
        self.context['result']['items'].append({'order_item_id': 11, 'decision': 'approved', 'quantity': 5,
                                              'unit': 'шт', 'selected': {'name': 'Исключённый товар'}, 'proposed_unit_price': 400})
        self.fetch.return_value = copy.deepcopy(self.order)
        self.invoice.side_effect = build_invoice_xlsx
        post.side_effect = [self.upload(), self.response({'id': 'offer-edited'})]
        response = self.client.post(self.url, data={**self.form, 'qty_10': '3', 'price_10': '125,50',
                                                   'provider_name_10': 'Исправленный товар', 'unit_10': 'упак'})
        self.assertEqual(response.status_code, 200, response.text)
        workbook = load_workbook(BytesIO(post.call_args_list[0].kwargs['files']['file'][1]))
        sheet = workbook.active
        self.assertEqual(sheet['A5'].value, 'Счёт на оплату № 278 от 22.10.2026')
        invoice_row = next(row for row in sheet if len(row) >= 6 and row[1].value == 'Исправленный товар')
        self.assertEqual([cell.value for cell in invoice_row[1:6]],
                         ['Исправленный товар', 3, 'упак', 125.5, 376.5])
        text = ' '.join(str(c.value or '') for row in sheet for c in row)
        self.assertNotIn('Исключённый товар', text)
        payload = post.call_args_list[1].kwargs['json']
        self.assertEqual(len(payload['offerItems']), 1)
        self.assertEqual(float(payload['offerItems'][0]['amount']), invoice_row[5].value)
        self.assertEqual(payload['producerOfferDate'], '2026-10-22')
        # The saved comparison remains the source for later review.
        self.assertEqual(self.context['result']['items'][0]['quantity'], 2)

    @patch('offer_panel_safe.requests.post')
    def test_selected_company_is_retained_when_order_is_unavailable(self, post):
        self.fetch.return_value = None
        response = self.client.post(self.url, data={**self.form, 'seller_id': 'sv', 'destination_account_id': 'sv-bank'})
        self.assertEqual(response.status_code, 409)
        self.save_settings.assert_called_once_with(456, self.sv, 'sv-bank', 0.22)
        self.assertIn('/dashboard/automation/jobs/456/invoice.xlsx', response.text)
        self.invoice.assert_not_called()
        post.assert_not_called()

    @patch('offer_panel_safe.requests.post')
    def test_mismatched_company_or_bank_is_rejected_before_upload(self, post):
        for changes in ({'seller_id': 'sv'}, {'seller_id': 'avior', 'destination_account_id': 'sv-bank'}):
            with self.subTest(changes=changes):
                response = self.client.post(self.url, data={**self.form, **changes})
                self.assertEqual(response.status_code, 409)
        post.assert_not_called()
        self.invoice.assert_not_called()
        self.save_settings.assert_not_called()

    @patch('offer_panel_safe.requests.post')
    def test_old_form_bank_selection_resolves_sv_without_defaulting_to_avior(self, post):
        post.side_effect = [self.upload(), self.response({'id': 'offer-sv'})]
        response = self.client.post(self.url, data={**self.form, 'destination_account_id': 'sv-bank'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.invoice.call_args.args[0]['seller']['inn'], self.sv['inn'])

    @patch('offer_panel_safe.requests.post')
    def test_save_firm_is_local_and_never_uploads_or_sends_offer(self, post):
        response = self.client.post('/dashboard/order/37463885/offer/settings', data={
            'seller_id': 'sv', 'destination_account_id': 'sv-bank', 'vat_rate': '0.2'}, follow_redirects=False)
        self.assertEqual(response.status_code, 303)
        self.save_settings.assert_called_once_with(456, self.sv, 'sv-bank', 0.2)
        post.assert_not_called()
        self.fetch.assert_not_called()

    def test_form_keeps_saved_sv_and_does_not_use_global_avior_account(self):
        self.context['result'].update(seller=self.sv, destination_account_id='sv-bank', vat_rate=0.2)
        with patch.dict('os.environ', {'ZAKUPAY_DESTINATION_ACCOUNT_ID': 'bank'}):
            response = self.client.get('/dashboard/order/37463885/offer')
        self.assertIn("value='saved' selected", response.text)
        self.assertIn("value='sv-bank' selected", response.text)
        self.assertIn("value='0.2' selected", response.text)
        self.assertNotIn('Банковский счёт ООО «АВИОР»', response.text)

    @patch("offer_panel_safe.requests.post")
    def test_missing_live_order_has_recovery_page_and_never_uploads(self, post):
        self.fetch.return_value = None
        response = self.client.post(self.url, data=self.form)
        self.assertEqual(response.status_code, 409)
        self.assertIn("не подтверждает её удаление", response.text)
        self.assertIn("/dashboard/automation/jobs/456/invoice.xlsx", response.text)
        self.invoice.assert_not_called()
        self.claim.assert_not_called()
        post.assert_not_called()

    @patch("offer_panel_safe.requests.post")
    def test_verification_is_read_only_even_when_order_is_available(self, post):
        response = self.client.get("/dashboard/order/37463885/offer/verify")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Заявка доступна в Закупай", response.text)
        self.fetch.assert_called_once_with(37463885, force=True)
        self.invoice.assert_not_called()
        self.claim.assert_not_called()
        post.assert_not_called()

    @patch("offer_panel_safe.requests.post")
    def test_verification_does_not_substitute_saved_order_when_live_order_is_absent(self, post):
        self.fetch.return_value = None
        response = self.client.get("/dashboard/order/37463885/offer/verify")
        self.assertEqual(response.status_code, 409)
        self.assertIn("В ответе Закупай нет заявки", response.text)
        self.invoice.assert_not_called()
        self.claim.assert_not_called()
        post.assert_not_called()

    @patch("offer_panel_safe.requests.post")
    def test_verification_detects_changed_order(self, post):
        self.fetch.return_value = {**self.order, "orderItems": [{"id": 10, "goodName": "Товар", "count": 3}]}
        response = self.client.get("/dashboard/order/37463885/offer/verify")
        self.assertEqual(response.status_code, 409)
        self.assertIn("изменились", response.text)
        self.invoice.assert_not_called()
        self.claim.assert_not_called()
        post.assert_not_called()


if __name__ == "__main__":
    unittest.main()
