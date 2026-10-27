"""Seller choices survive reload/search and agree with downloaded invoices."""
import copy
import gc
import json
import os
import tempfile
import unittest
from io import BytesIO
from unittest.mock import Mock, patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import automation_pipeline as pipeline
from invoice_generator import SELLER


class SavedInvoiceSellerTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        # sqlite connection contexts commit but close only when collected on Windows.
        self.addCleanup(gc.collect)
        db = patch.multiple(pipeline, DB_PATH=os.path.join(self.tmp.name, 'jobs.db'), DATABASE_URL='', _initialized_database_key=None)
        db.start()
        self.addCleanup(db.stop)
        self.sv = {**SELLER, 'name': 'ООО «СВ»', 'inn': '5638073633', 'kpp': '561101001', 'account': '4' * 20}
        self.order = {'id': 42, 'orderItems': [{'id': 10, 'goodName': 'Товар', 'count': 2}]}
        self.result = {'order_id': 42, 'invoice_number': 278, 'customer': {'name': 'ООО Покупатель', 'inn': '7716997861', 'kpp': '770801001'},
                       'items': [{'position': 1, 'order_item_id': 10, 'decision': 'approved', 'selected': {'name': 'Товар'},
                                  'requested_name': 'Товар', 'quantity': 2, 'unit': 'шт', 'proposed_unit_price': 100}],
                       'summary': {'positions': 1, 'auto_ready': 1}, 'status': 'ready_for_review'}
        with pipeline._connect() as conn:
            self.job = conn.execute('INSERT INTO automation_jobs(dedupe_key,order_id,event_type,status,order_json,result_json,created_at,updated_at,invoice_number) VALUES(?,?,?,?,?,?,?,?,?)',
                                    ('seller-test', 42, 'api_order', 'ready_for_review', json.dumps(self.order), json.dumps(self.result), 'now', 'now', 278)).lastrowid
        app = FastAPI()
        pipeline.install_automation_pipeline(app, Mock())
        self.client = TestClient(app)
        self.addCleanup(self.client.close)

    def test_saved_sv_is_used_by_both_download_routes_after_reload(self):
        pipeline.save_invoice_settings(self.job, self.sv, 'sv-account', 0.2)
        context = pipeline.load_automation_offer_context(42)
        self.assertEqual(context['result']['seller'], self.sv)
        self.assertEqual(context['result']['destination_account_id'], 'sv-account')
        for prefix in ('/automation', '/dashboard/automation'):
            response = self.client.get(f'{prefix}/jobs/{self.job}/invoice.xlsx')
            self.assertEqual(response.status_code, 200, response.text if response.status_code != 200 else '')
            sheet = load_workbook(BytesIO(response.content)).active
            self.assertIn(self.sv['inn'], sheet['A7'].value)
            self.assertNotIn(SELLER['inn'], sheet['A7'].value)
            self.assertIn(self.sv['inn'], response.headers['content-disposition'])

    def test_supplier_search_does_not_reset_selected_seller_or_tax(self):
        pipeline.save_invoice_settings(self.job, self.sv, 'sv-account', 0.2)
        with patch.object(pipeline, 'build_vi_draft', return_value=copy.deepcopy(self.result)):
            response = self.client.post(f'/dashboard/automation/jobs/{self.job}/refresh', follow_redirects=False)
        self.assertEqual(response.status_code, 303)
        saved = pipeline.load_automation_offer_context(42)['result']
        self.assertEqual(saved['seller'], self.sv)
        self.assertEqual(saved['vat_rate'], 0.2)
        self.assertEqual(saved['destination_account_id'], 'sv-account')

    def test_existing_tile_quote_seller_is_recovered_only_for_same_job(self):
        with pipeline._connect() as conn:
            conn.execute('INSERT INTO tile_quotes(job_id,catalog_id,snapshot_json,created_at) VALUES(?,?,?,?)',
                         (self.job + 1, 1, json.dumps({'seller': self.sv, 'vat_rate': 0.2}), 'now'))
        self.assertNotIn('seller', pipeline.load_automation_offer_context(42)['result'])
        with pipeline._connect() as conn:
            conn.execute('INSERT INTO tile_quotes(job_id,catalog_id,snapshot_json,created_at) VALUES(?,?,?,?)',
                         (self.job, 1, json.dumps({'seller': self.sv, 'vat_rate': 0.2}), 'now'))
        response = self.client.get(f'/dashboard/automation/jobs/{self.job}/invoice.xlsx')
        self.assertEqual(response.status_code, 200)
        self.assertIn(self.sv['inn'], load_workbook(BytesIO(response.content)).active['A7'].value)

    def test_seller_change_is_blocked_for_sent_and_uncertain_offers(self):
        pipeline.save_invoice_settings(self.job, self.sv, 'sv-account', 0.2)
        pipeline.begin_offer_submission(self.job, 42, 'attempt', 'guid', 'hash', 'payload', '278', 'file', {})
        with self.assertRaisesRegex(ValueError, 'заблокировано'):
            pipeline.save_invoice_settings(self.job, SELLER, 'avior-account', 0.22)
        self.assertEqual(pipeline.load_automation_offer_context(42)['result']['seller']['inn'], self.sv['inn'])

    def test_invoice_download_requires_a_selected_company(self):
        for prefix in ('/automation', '/dashboard/automation'):
            response = self.client.get(f'{prefix}/jobs/{self.job}/invoice.xlsx')
            self.assertEqual(response.status_code, 409)
            self.assertIn('Выберите нашу фирму', response.json()['detail'])
        review = self.client.get(f'/dashboard/automation/jobs/{self.job}/review')
        self.assertIn("name='invoice_seller_id' required", review.text)
        self.assertIn('Сформировать и скачать счёт', review.text)
        self.assertNotIn("value='avior' selected", review.text)

    def test_generate_saves_company_and_rows_before_downloading(self):
        from porcelain import save_seller
        key = save_seller(self.sv)
        pipeline.save_invoice_settings(self.job, SELLER, 'old-avior-account', 0.22)
        response = self.client.post(f'/dashboard/automation/jobs/{self.job}/review', data={
            'invoice_seller_id': key, 'invoice_vat_rate': '0', 'review_action': 'download',
            'include_1': '1', 'quantity_1': '3', 'price_1': '125.5'})
        self.assertEqual(response.status_code, 200, response.text if response.status_code != 200 else '')
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertIn(self.sv['inn'], sheet['A7'].value)
        self.assertNotIn(SELLER['inn'], sheet['A7'].value)
        cells = ' '.join(str(cell.value or '') for row in sheet for cell in row)
        self.assertIn('Без НДС', cells)
        item = next(row for row in sheet if row[1].value == 'Товар')
        self.assertEqual([cell.value for cell in item[2:6]], [3, 'шт', 125.5, 376.5])
        saved = pipeline.load_automation_offer_context(42)['result']
        self.assertEqual(saved['seller'], self.sv)
        self.assertEqual(saved['vat_rate'], 0)
        self.assertEqual(saved['destination_account_id'], '')
        review = self.client.get(f'/dashboard/automation/jobs/{self.job}/review')
        self.assertIn("value='saved' selected", review.text)
        self.assertIn("value='0' selected", review.text)

    def test_invalid_company_or_tax_does_not_modify_the_invoice(self):
        for changes in ({'invoice_seller_id': ''}, {'invoice_seller_id': 'deleted-firm'},
                        {'invoice_seller_id': 'avior', 'invoice_vat_rate': '0.3'}):
            with self.subTest(changes=changes):
                response = self.client.post(f'/dashboard/automation/jobs/{self.job}/review', data={
                    'include_1': '1', 'quantity_1': '99', 'price_1': '1', **changes})
                self.assertEqual(response.status_code, 409)
                saved = pipeline.load_automation_offer_context(42)['result']
                self.assertNotIn('seller', saved)
                self.assertEqual(saved['items'][0]['quantity'], 2)

    def test_review_cannot_change_company_during_submission(self):
        pipeline.save_invoice_settings(self.job, self.sv, 'sv-account', 0.2)
        pipeline.begin_offer_submission(self.job, 42, 'attempt', 'guid', 'hash', 'payload', '278', 'file', {})
        response = self.client.post(f'/dashboard/automation/jobs/{self.job}/review', data={
            'invoice_seller_id': 'avior', 'invoice_vat_rate': '0.22',
            'include_1': '1', 'quantity_1': '2', 'price_1': '100'})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(pipeline.load_automation_offer_context(42)['result']['seller'], self.sv)

    def test_offer_confirmation_receives_the_company_selected_above_the_table(self):
        from porcelain import save_seller
        key = save_seller(self.sv)
        response = self.client.post(f'/dashboard/automation/jobs/{self.job}/review', data={
            'invoice_seller_id': key, 'invoice_vat_rate': '0.2', 'review_action': 'offer',
            'include_1': '1', 'quantity_1': '2', 'price_1': '100'}, follow_redirects=False)
        self.assertEqual(response.status_code, 303)
        self.assertEqual(response.headers['location'], '/dashboard/order/42/offer')
        self.assertEqual(pipeline.load_automation_offer_context(42)['result']['seller'], self.sv)

    def test_saved_successful_offer_id_is_recovered_without_resending(self):
        from offer_responses import UNCONFIRMED_OFFER_ERROR
        claim = pipeline.begin_offer_submission(self.job, 42, 'attempt', 'guid', 'hash', 'payload', '278', 'file', {})
        reply = {'result': True, 'offerId': 340932504,
                 'offerUrl': 'https://zakupay.sel-be.ru/core/offers/340932504/report'}
        pipeline.finish_offer_submission(claim['id'], 'unknown', response=reply, error=UNCONFIRMED_OFFER_ERROR)
        with patch('requests.sessions.Session.request') as remote:
            self.assertEqual(pipeline.reconcile_offer_confirmations(), [claim['id']])
            self.assertEqual(pipeline.reconcile_offer_confirmations(), [])
            remote.assert_not_called()
        record = pipeline.latest_offer_submission(self.job)
        self.assertEqual(record['status'], 'confirmed')
        self.assertEqual(record['offer_id'], '340932504')
        result = pipeline.load_automation_offer_context(42)['result']
        self.assertTrue(result['live_offer_created'])
        self.assertEqual(result['live_offer_id'], '340932504')
        self.assertEqual(result['live_offer_file_id'], 'file')

    def test_recovery_keeps_ambiguous_or_failed_responses_blocked(self):
        from offer_responses import UNCONFIRMED_OFFER_ERROR
        claim = pipeline.begin_offer_submission(self.job, 42, 'attempt', 'guid', 'hash', 'payload', '278', 'file', {})
        cases = [(None, 'Timeout'), ({'result': True, 'offerId': 340932504}, 'HTTP 502'),
                 ({'result': False, 'offerId': 340932504}, UNCONFIRMED_OFFER_ERROR),
                 ({'result': True, 'id': 'file-id'}, UNCONFIRMED_OFFER_ERROR),
                 ({'offerId': 340932504}, UNCONFIRMED_OFFER_ERROR)]
        for response, error in cases:
            with self.subTest(response=response, error=error):
                with pipeline._connect() as conn:
                    conn.execute("UPDATE offer_submissions SET status='unknown', response_json=?, error=? WHERE id=?",
                                 (json.dumps(response) if response is not None else None, error, claim['id']))
                self.assertEqual(pipeline.reconcile_offer_confirmations(), [])
                self.assertEqual(pipeline.latest_offer_submission(self.job)['status'], 'unknown')
                self.assertFalse(pipeline.load_automation_offer_context(42)['result'].get('live_offer_created'))

    def default_choice(self, enabled='1', seller=None, vat='0.2'):
        from porcelain import save_seller
        key = save_seller(seller or self.sv)
        return self.client.post(f'/dashboard/automation/jobs/{self.job}/seller-default',
            headers={'X-Requested-With': 'invoice-preferences'},
            data={'enabled': enabled, 'seller_id': key, 'vat_rate': vat})

    def test_default_firm_survives_reload_and_is_used_in_new_invoice(self):
        self.assertEqual(self.default_choice().status_code, 200)
        page = self.client.get(f'/dashboard/automation/jobs/{self.job}/review')
        self.assertIn("id='invoice-default' type='checkbox' checked", page.text)
        self.assertIn("id='invoice-seller-options' hidden", page.text)
        self.assertIn(f"data-save-url='/dashboard/automation/jobs/{self.job}/seller-default'", page.text)
        with pipeline._connect() as conn:
            original = conn.execute('SELECT result_json FROM automation_jobs WHERE id=?', (self.job,)).fetchone()[0]
        self.assertNotIn('seller', json.loads(original))
        # A second request without an explicit seller adopts the durable preference on save.
        order = copy.deepcopy(self.order)
        order['id'] = 43
        result = copy.deepcopy(self.result)
        result['order_id'] = 43
        with pipeline._connect() as conn:
            new_job = conn.execute('INSERT INTO automation_jobs(dedupe_key,order_id,event_type,status,order_json,result_json,created_at,updated_at,invoice_number) VALUES(?,?,?,?,?,?,?,?,?)',
                ('new-default', 43, 'api_order', 'ready_for_review', json.dumps(order), json.dumps(result), 'now', 'now', 279)).lastrowid
        response = self.client.post(f'/dashboard/automation/jobs/{new_job}/review', data={
            'review_action': 'download', 'include_1': '1', 'quantity_1': '2', 'price_1': '100'})
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        cells = ' '.join(str(c.value or '') for r in workbook.active for c in r)
        self.assertIn(self.sv['inn'], cells)
        self.assertNotIn(SELLER['inn'], cells)
        self.assertIn('20%', cells)
        self.assertEqual(pipeline.load_automation_offer_context(43)['result']['seller'], self.sv)
        self.assertEqual(pipeline.load_automation_offer_context(43)['result']['vat_rate'], 0.2)

    def test_default_survives_database_reinitialization_and_new_authenticated_session(self):
        import security
        self.assertEqual(self.default_choice().status_code, 200)
        pipeline._initialized_database_key = None
        app = FastAPI()
        pipeline.install_automation_pipeline(app, Mock())
        security.install_security(app)
        with TestClient(app) as client:
            url = f'/dashboard/automation/jobs/{self.job}/review'
            self.assertEqual(client.get(url, follow_redirects=False).status_code, 303)
            client.cookies.set(security.SESSION_COOKIE, security._session_token())
            page = client.get(url)
            self.assertEqual(page.status_code, 200)
            self.assertIn("id='invoice-seller-options' hidden", page.text)
            self.assertIn(self.sv['inn'], page.text)
            self.assertEqual(client.post(f'/dashboard/automation/jobs/{self.job}/seller-default',
                headers={'X-Requested-With': 'invoice-preferences'},
                data={'enabled': '1', 'seller_id': 'saved', 'vat_rate': '0.2'}).status_code, 200)

    def test_default_does_not_replace_existing_saved_or_sent_invoice(self):
        pipeline.save_invoice_settings(self.job, SELLER, 'avior-account', 0.22)
        with pipeline._connect() as conn:
            before = conn.execute('SELECT result_json FROM automation_jobs WHERE id=?', (self.job,)).fetchone()[0]
        self.assertEqual(self.default_choice().status_code, 200)
        page = self.client.get(f'/dashboard/automation/jobs/{self.job}/review')
        self.assertNotIn("id='invoice-seller-options' hidden", page.text)
        response = self.client.get(f'/dashboard/automation/jobs/{self.job}/invoice.xlsx')
        self.assertIn(SELLER['inn'], load_workbook(BytesIO(response.content)).active['A7'].value)
        with pipeline._connect() as conn:
            self.assertEqual(before, conn.execute('SELECT result_json FROM automation_jobs WHERE id=?', (self.job,)).fetchone()[0])
            sent = json.loads(before)
            sent['live_offer_created'] = True
            conn.execute('UPDATE automation_jobs SET result_json=?, status=? WHERE id=?', (json.dumps(sent), 'offer_created', self.job))
        self.client.get(f'/dashboard/automation/jobs/{self.job}/review')
        self.assertEqual(pipeline.load_automation_offer_context(42)['result']['seller'], SELLER)

    def test_default_can_be_disabled_and_replaced_and_is_account_scoped(self):
        self.assertEqual(self.default_choice().status_code, 200)
        from invoice_preferences import load
        with pipeline._connect() as conn:
            self.assertIsNone(load(conn, pipeline._execute, 'another-user'))
        self.assertEqual(self.default_choice(enabled='0').status_code, 200)
        page = self.client.get(f'/dashboard/automation/jobs/{self.job}/review')
        self.assertNotIn("id='invoice-seller-options' hidden", page.text)
        self.assertEqual(self.default_choice(seller=SELLER, vat='0').status_code, 200)
        page = self.client.get(f'/dashboard/automation/jobs/{self.job}/review')
        self.assertIn("id='invoice-seller-options' hidden", page.text)
        self.assertIn("value='0' selected", page.text)

    def test_unpin_keeps_current_choice_available_for_invoice_save(self):
        self.assertEqual(self.default_choice().status_code, 200)
        self.assertEqual(self.default_choice(enabled='0').status_code, 200)
        response = self.client.post(f'/dashboard/automation/jobs/{self.job}/review', data={
            'invoice_seller_id': 'saved', 'invoice_vat_rate': '0.2', 'review_action': 'download',
            'include_1': '1', 'quantity_1': '2', 'price_1': '100'})
        self.assertEqual(response.status_code, 200)
        self.assertIn(self.sv['inn'], load_workbook(BytesIO(response.content)).active['A7'].value)

    def test_invalid_or_cross_site_default_update_leaves_preference_unchanged(self):
        self.assertEqual(self.default_choice().status_code, 200)
        self.assertEqual(self.default_choice(vat='nan').status_code, 422)
        self.assertEqual(self.client.post(f'/dashboard/automation/jobs/{self.job}/seller-default', data={'enabled': '0'}).status_code, 403)
        page = self.client.get(f'/dashboard/automation/jobs/{self.job}/review')
        self.assertIn("value='0.2' selected", page.text)
        self.assertIn("id='invoice-seller-options' hidden", page.text)


if __name__ == '__main__':
    unittest.main()
