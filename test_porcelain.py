import copy
import io
import json
import os
import tempfile
import unittest
import zipfile
from unittest.mock import Mock, patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import automation_pipeline as pipeline
import porcelain as tiles
import porcelain_panel as panel
import security
from invoice_generator import SELLER


CUSTOMER = {'name':'ООО Покупатель','inn':'7716997861','kpp':'770801001'}
PRODUCT = {'id':'KH-GRAY','name':'Наш серый мрамор','format':'60x60','color':'Серый','pattern':'Мрамор',
           'surface':'Матовая','thickness_mm':9,'stock_m2':100,'box_m2':1.44,'pallet_m2':103.68,
           'kg_m2_confirmed':22,'image':'assets/KH-GRAY.jpg','price_rub_m2':None,'vat':'Не указан'}
ORDER = {'id':123,'name':'Керамогранит в Москву','deliveryAddress':{'fullAddress':'Москва, Тестовая, 1'},'customer':CUSTOMER,
         'orderItems':[{'id':101,'goodName':'Керамогранит серый под мрамор 600х600х9 мм матовый', 'count':40,'unit':{'name':'м²'}}]}


def archive(products=None, image=b'\xff\xd8test-photo\xff\xd9'):
    data={'schema_version':'1.1-draft','stock_as_of':'2026-10-17T13:40:00+03:00','products':products or [copy.deepcopy(PRODUCT)],
          'shipping':{'scope':'Москва и МО','tariffs':[{'vehicle':'Газель 1,5 т','capacity_kg':1500,
          'low':4300,'mid':4790,'high':5800,'km_low':35,'km_mid':55,'km_high':75,'center':0}]}}
    output=io.BytesIO()
    with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as z:
        z.writestr('catalog.json',json.dumps(data,ensure_ascii=False))
        for p in data['products']:
            if p.get('image') and p['image'] not in z.namelist():z.writestr(p['image'],image)
    return output.getvalue()


class PorcelainTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.settings=patch.multiple(pipeline,DB_PATH=os.path.join(self.tmp.name,'tiles.db'),DATABASE_URL='',_initialized_database_key=None)
        self.settings.start();self.addCleanup(self.settings.stop);self.addCleanup(self.tmp.cleanup)
        self.fetch=Mock(return_value=[copy.deepcopy(ORDER)])
        app=FastAPI();pipeline.install_automation_pipeline(app,Mock(side_effect=AssertionError('No order detail lookup')))
        panel.install_porcelain(app,self.fetch);security.install_security(app)
        self.auth=patch.object(security,'_valid_session',return_value=True);self.auth.start();self.addCleanup(self.auth.stop)
        self.client=TestClient(app);self.addCleanup(self.client.close)
        self.catalog_id=tiles.import_catalog(archive());self.catalog=tiles.load_catalog(self.catalog_id)
        self.job=pipeline.save_api_order_for_manual_start(copy.deepcopy(ORDER))['job_id']

    def ready_form(self):
        return {'product_1':'KH-GRAY','area_1':'40','reserve_1':'0','price_1':'1000','approve_1':'yes',
                'address':'Москва, Тестовая, 1','zone':'mkad','unloading':'0','vehicle':'Газель 1,5 т',
                'transport_confirmed':'yes','seller':'avior','seller_confirmed':'yes','vat':'22','invoice_number':'К-1'}

    def test_domain_matching_checks_size_color_design_surface_and_unknown_fields(self):
        check=tiles.assess(ORDER['orderItems'][0]['goodName'],PRODUCT)
        self.assertFalse(check['conflicts']);self.assertIn('Формат совпадает',check['reasons'])
        for text in ['Керамогранит белый 600х600х9 мм', 'Керамогранит 60х120 см',
                     'Керамогранит серый под дерево', 'Керамогранит 600х600х20 мм','Керамогранит полированный']:
            with self.subTest(text=text):self.assertTrue(tiles.assess(text,PRODUCT)['conflicts'])
        p={**PRODUCT,'thickness_mm':None}
        self.assertIn('Толщина не указан в каталоге',tiles.assess(ORDER['orderItems'][0]['goodName'],p)['unknown'])
        self.assertTrue(tiles.is_tile('Керамогранит 60x60'))
        self.assertFalse(tiles.is_tile('Клей для керамогранита'))
        self.assertFalse(tiles.is_tile('Диск для резки керамогранита'))

    def test_boxes_weight_delivery_rounding_and_summed_stock(self):
        quote=tiles.make_quote(ORDER,self.catalog,self.ready_form())
        self.assertEqual(quote['issues'],[])
        line=quote['lines'][0]
        self.assertEqual(line['boxes'],28);self.assertEqual(line['quantity'],40.32)
        self.assertAlmostEqual(line['weight_kg'],917.04)
        self.assertEqual(quote['shipping']['mid'],4790);self.assertEqual(quote['delivery_price'],5000)
        self.assertEqual(quote['total'],45320)
        order=copy.deepcopy(ORDER);order['orderItems']*=2
        form=self.ready_form();form.update({'product_2':'KH-GRAY','area_2':'65','price_2':'1000','approve_2':'yes'})
        quote=tiles.make_quote(order,self.catalog,form)
        self.assertTrue(any('по сумме всех строк' in issue for issue in quote['issues']))

    def test_missing_price_unknown_stock_and_missing_transport_never_create_invoice(self):
        form=self.ready_form();form['price_1']='';form['unloading']='';form['transport_confirmed']=''
        quote=tiles.make_quote(ORDER,self.catalog,form)
        self.assertIsNone(quote['total']);self.assertIsNone(quote['shipping'])
        with self.assertRaises(ValueError):tiles.invoice_draft(quote)
        unknown=copy.deepcopy(self.catalog);unknown['data']['products'][0]['stock_m2']=None
        quote=tiles.make_quote(ORDER,unknown,self.ready_form())
        self.assertTrue(any('Наличие неизвестно' in error for error in quote['issues']))
        form=self.ready_form();form['zone']='Магадан'
        self.assertIsNone(tiles.make_quote(ORDER,self.catalog,form)['shipping'])

    def test_pages_and_sync_do_not_run_supplier_search_or_send_anything(self):
        with patch.object(pipeline,'build_vi_draft',side_effect=AssertionError('No supplier lookups')):
            self.assertEqual(self.client.get(panel.URL).status_code,200)
            response=self.client.get(f'/dashboard/automation/jobs/{self.job}/tiles')
            self.assertIn('Наш серый мрамор',response.text)
            self.assertIn('Москва, Тестовая, 1',response.text)
            self.assertIn('area_1',response.text)
            self.assertIn("name='area_1' value='40'",response.text)
            self.fetch.assert_not_called()
            self.assertEqual(self.client.post(panel.URL+'/sync',follow_redirects=False).status_code,303)
            self.fetch.assert_called_once_with(force=True)

    def test_prepare_edit_and_bundle_snapshot_isolation_and_customer_files(self):
        response=self.client.post(f'/dashboard/automation/jobs/{self.job}/tiles/prepare',
            data={'catalog_id':self.catalog_id,'product_1':'KH-GRAY','area_1':'40'},follow_redirects=False)
        self.assertEqual(response.status_code,303)
        url=response.headers['location'];quote_id=int(url.rsplit('/',1)[1])
        self.assertEqual(self.client.get(url+'/bundle').status_code,409)
        response=self.client.post(url+'/revise',data=self.ready_form(),follow_redirects=False)
        ready_url=response.headers['location'];ready_id=int(ready_url.rsplit('/',1)[1])
        self.assertEqual(tiles.load_quote(quote_id)['snapshot']['lines'][0]['proposed_unit_price'],None)
        # A later catalogue import cannot swap the selected photograph or price.
        tiles.import_catalog(archive([{**PRODUCT,'price_rub_m2':7777}],image=b'NEW-IMAGE'))
        packet=self.client.get(ready_url+'/bundle')
        self.assertEqual(packet.status_code,200)
        with zipfile.ZipFile(io.BytesIO(packet.content)) as z:
            self.assertEqual(set(z.namelist()),{'invoice.xlsx','proposal.html','photos/KH-GRAY.jpg'})
            self.assertNotIn(b'NEW-IMAGE',z.read('photos/KH-GRAY.jpg'))
            self.assertNotIn('Грузоподъёмность',z.read('proposal.html').decode())
            self.assertNotIn('4790',z.read('proposal.html').decode())
            book=load_workbook(io.BytesIO(z.read('invoice.xlsx')),data_only=True)
            values=[c.value for row in book.active for c in row]
            self.assertIn(45320,values);self.assertIn('В том числе НДС 22%:',values)
        self.assertEqual(tiles.load_quote(ready_id)['snapshot']['seller']['inn'],SELLER['inn'])
        history=self.client.get(f'/dashboard/automation/jobs/{self.job}/tiles')
        self.assertIn(ready_url,history.text)

    def test_clearing_auto_delivery_recalculates_after_quantity_change(self):
        form=self.ready_form();form['sale_delivery']=''
        first=tiles.make_quote(ORDER,self.catalog,form)
        self.assertEqual(first['delivery_price'],5000)
        form['area_1']='90'
        second=tiles.make_quote(ORDER,self.catalog,form)
        self.assertEqual(second['shipping']['trips'],2)
        self.assertEqual(second['delivery_price'],10000)
        key=tiles.save_quote(self.job,self.catalog,second)
        response=self.client.get(f'{panel.URL}/quotes/{key}')
        self.assertIn("name='sale_delivery' value=''",response.text)

    def test_buyer_can_be_completed_without_replacing_original_request(self):
        order=copy.deepcopy(ORDER);order['customer']={'fullName':'Неизвестный покупатель'}
        form=self.ready_form();form.update(customer_name=CUSTOMER['name'],customer_inn=CUSTOMER['inn'],customer_kpp=CUSTOMER['kpp'])
        quote=tiles.make_quote(order,self.catalog,form)
        self.assertFalse(quote['issues'])
        self.assertEqual(quote['customer']['name'],CUSTOMER['name'])
        self.assertEqual(order['customer'],{'fullName':'Неизвестный покупатель'})

    def test_select_other_own_firm_and_no_vat_invoice(self):
        firm={**SELLER,'name':'ИП Тестовый','inn':'123456789012','kpp':'','account':'40702810000000000001'}
        key=tiles.save_seller(firm);form=self.ready_form();form.update(seller=key,vat='0')
        quote=tiles.make_quote(ORDER,self.catalog,form)
        packet=panel.bundle({'snapshot':quote,'catalog_id':self.catalog_id})
        with zipfile.ZipFile(io.BytesIO(packet)) as z:
            book=load_workbook(io.BytesIO(z.read('invoice.xlsx')),data_only=True)
            values=[str(c.value or '') for row in book.active for c in row]
            self.assertIn('Без НДС:',values)
            self.assertTrue(any('ИП Тестовый' in v for v in values))
            self.assertFalse(any('5017134957' in v for v in values))

    def test_archive_rejects_missing_image_and_duplicate_ids(self):
        with self.assertRaises(ValueError):tiles.import_catalog(archive([PRODUCT,PRODUCT]))
        output=io.BytesIO()
        with zipfile.ZipFile(output,'w') as z:z.writestr('catalog.json',json.dumps({'products':[PRODUCT]}))
        with self.assertRaises(ValueError):tiles.import_catalog(output.getvalue())


if __name__=='__main__':unittest.main()
