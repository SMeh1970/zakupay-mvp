"""Generate a reviewable invoice using the selected seller's requisites."""

from __future__ import annotations

from io import BytesIO
from datetime import date
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def customer_product_link(value):
    """Accept only our public product pages, never supplier or arbitrary URLs."""
    value = str(value or '')
    return value if re.fullmatch(r'https://keramahall-delivery\.mehmans\.chatgpt\.site/products/[A-Za-z0-9_-]{1,80}', value) else None


def product_qr_png(url):
    import qrcode
    if not customer_product_link(url):
        raise ValueError('Недопустимая ссылка на товар')
    output = BytesIO()
    qrcode.make(url, error_correction=qrcode.constants.ERROR_CORRECT_M, box_size=6, border=4).save(output, format='PNG')
    return output.getvalue()


SELLER = {
    "name": 'ООО «АВИОР»',
    "full_name": 'Общество с ограниченной ответственностью «АВИОР»',
    "inn": "5017134957",
    "kpp": "501701001",
    "ogrn": "1235000160713",
    "address": "143600, Московская область, Волоколамский г.о., г. Волоколамск, ул. Панфилова, д. 5, часть 1/2, помещ. 24",
    "bank": 'ДО «Даниловский (ЮЛ)» в г. Москва АО «АЛЬФА-БАНК»',
    "account": "40702810502860023762",
    "bik": "044525593",
    "corr_account": "30101810200000000593",
    "director": "Баталова Елена Геннадьевна",
}

AVAILABILITY_CONFIRMATION_NOTE = (
    "Перед согласованием счёта просим запросить у нас подтверждение актуального наличия товаров."
)


def validate_seller(profile):
    if not isinstance(profile, dict):
        raise ValueError('Выберите реквизиты поставщика счёта')
    result = {key: ' '.join(str(profile.get(key) or '').split()) for key in SELLER}
    for key in ('name', 'address', 'bank', 'director'):
        if not result[key]:
            raise ValueError('Заполните реквизиты нашей фирмы: ' + key)
    for key, lengths in (('inn', {10, 12}), ('bik', {9}), ('account', {20}), ('corr_account', {20})):
        if not result[key].isascii() or not result[key].isdigit() or len(result[key]) not in lengths:
            raise ValueError('Проверьте реквизиты нашей фирмы: ' + key)
    if len(result['inn']) == 10 and (not result['kpp'].isascii() or not result['kpp'].isdigit() or len(result['kpp']) != 9):
        raise ValueError('Укажите КПП нашей организации')
    return result


def invoice_filename(draft):
    seller = validate_seller(draft['seller']) if 'seller' in draft else SELLER
    number = re.sub(r'[^A-Za-z0-9_-]', '_', str(draft.get('invoice_number') or 'draft'))[:60]
    order = re.sub(r'[^A-Za-z0-9_-]', '_', str(draft.get('order_id') or ''))[:60]
    return f"invoice_{seller['inn']}_{number}_order_{order}.xlsx"


def normalize_customer(customer: dict | None) -> dict:
    customer = customer if isinstance(customer, dict) else {}
    def clean(value):
        return " ".join(str(value or "").split())
    address = customer.get("legalAddress") or customer.get("address")
    if isinstance(address, dict):
        address = address.get("fullAddress") or address.get("name")
    return {
        "name": clean(customer.get("fullName") or customer.get("name") or customer.get("shortName")),
        "inn": clean(customer.get("inn")),
        "kpp": clean(customer.get("kpp")),
        "legalAddress": clean(address),
    }


def customer_validation_error(customer: dict | None) -> str | None:
    customer = normalize_customer(customer)
    errors = []
    if not customer["name"] or customer["name"].casefold() in {"заказчик по заявке закупай", "—", "-"}:
        errors.append("наименование плательщика")
    if not re.fullmatch(r"(?:[0-9]{10}|[0-9]{12})", customer["inn"]):
        errors.append("ИНН плательщика (10 или 12 цифр)")
    if len(customer["inn"]) == 10 and not re.fullmatch(r"[0-9]{4}[0-9A-Z]{2}[0-9]{3}", customer["kpp"]):
        errors.append("КПП организации (9 знаков)")
    elif customer["kpp"] and not re.fullmatch(r"[0-9]{4}[0-9A-Z]{2}[0-9]{3}", customer["kpp"]):
        errors.append("корректный КПП")
    if errors:
        return "Счёт нельзя сформировать или отправить: укажите " + ", ".join(errors) + ". Сохраните реквизиты плательщика в проверке заявки."
    return None


def _customer_text(customer: dict | None) -> str:
    error = customer_validation_error(customer)
    if error:
        raise ValueError(error)
    customer = normalize_customer(customer)
    parts = [customer["name"], f"ИНН {customer['inn']}"]
    if customer["kpp"]:
        parts.append(f"КПП {customer['kpp']}")
    if customer["legalAddress"]:
        parts.append(customer["legalAddress"])
    return ", ".join(parts)


def build_invoice_xlsx(draft: dict) -> bytes:
    seller = validate_seller(draft['seller']) if 'seller' in draft else SELLER
    seller_kpp = f", КПП {seller['kpp']}" if seller.get('kpp') else ''
    vat_rate = float(draft.get('vat_rate', 0.22))
    if vat_rate not in {0, 0.1, 0.2, 0.22}:
        raise ValueError('Выберите ставку НДС для счёта')
    all_rows = draft.get("items") or []
    from supplier_comparison import selected_allowed
    rows = [row for row in all_rows if row.get("decision") in {"auto_ready", "approved"} and selected_allowed(row)]
    if not rows:
        raise ValueError("Счёт нельзя сформировать: нет ни одной подтверждённой позиции")
    customer_text = _customer_text(draft.get("customer"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Счёт"
    widths = [6, 68, 13, 12, 16, 18]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    bold = Font(bold=True)
    title = Font(bold=True, size=16)
    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="D9EAF7")

    ws.merge_cells("A1:F1")
    ws["A1"] = seller["bank"]
    ws["A1"].font = bold
    ws.merge_cells("A2:C2")
    ws["A2"] = f"БИК {seller['bik']}"
    ws.merge_cells("D2:F2")
    ws["D2"] = f"к/с {seller['corr_account']}"
    ws.merge_cells("A3:C3")
    ws["A3"] = f"Получатель: {seller['name']}, ИНН {seller['inn']}{seller_kpp}"
    ws.merge_cells("D3:F3")
    ws["D3"] = f"р/с {seller['account']}"

    number = draft.get("invoice_number")
    invoice_date = date.fromisoformat(str(draft["invoice_date"])) if draft.get("invoice_date") else date.today()
    ws.merge_cells("A5:F5")
    ws["A5"] = f"Счёт на оплату № {number} от {invoice_date.strftime('%d.%m.%Y')}"
    ws["A5"].font = title

    ws.merge_cells("A7:F7")
    ws["A7"] = f"Поставщик: {seller['name']}, ИНН {seller['inn']}{seller_kpp}, {seller['address']}"
    ws["A7"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A8:F8")
    ws["A8"] = f"Покупатель: {customer_text}"
    ws["A8"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[8].height = max(30, 15 * ((len(customer_text) + 11) // 120 + 1))
    ws.merge_cells("A9:F9")
    ws["A9"] = f"Основание: заявка Закупай № {draft.get('order_id')}"

    headers = ["№", "Наименование", "Кол-во", "Ед.", "Цена, руб.", "Сумма, руб."]
    start = 11
    for col, value in enumerate(headers, 1):
        cell = ws.cell(start, col, value)
        cell.font = bold
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    total = 0.0
    product_links = []
    for row_num, item in enumerate(rows, start + 1):
        qty = float(item.get("quantity") or 0)
        price = float(item.get("proposed_unit_price") or 0)
        amount = round(qty * price, 2)
        total += amount
        values = [
            row_num - start,
            (item.get("selected") or {}).get("name") or item.get("requested_name"),
            qty,
            item.get("unit") or "шт.",
            price,
            amount,
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row_num, col, value)
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row_num, 5).number_format = '#,##0.00'
        ws.cell(row_num, 6).number_format = '#,##0.00'
        product_url = customer_product_link(item.get('product_url'))
        if product_url:
            ws.cell(row_num, 2).hyperlink = product_url
            ws.cell(row_num, 2).font = Font(color='17569B', underline='single')
            product_links.append((str(values[1] or ''), product_url))

    footer = start + 1 + len(rows)
    vat = round(total * vat_rate / (1 + vat_rate), 2)
    vat_label = f"В том числе НДС {vat_rate * 100:g}%:" if vat_rate else 'Без НДС:'
    for label, value in (("Итого:", total), (vat_label, vat), ("Всего к оплате:", total)):
        ws.merge_cells(start_row=footer, start_column=1, end_row=footer, end_column=5)
        ws.cell(footer, 1, label).font = bold
        ws.cell(footer, 6, value).font = bold
        ws.cell(footer, 6).number_format = '#,##0.00'
        footer += 1

    ws.merge_cells(start_row=footer + 1, start_column=1, end_row=footer + 1, end_column=6)
    ws.cell(footer + 1, 1, draft.get('payment_note') or "Оплата: 100% предоплата. Доставка включена в стоимость.")
    ws.merge_cells(start_row=footer + 2, start_column=1, end_row=footer + 2, end_column=6)
    ws.cell(footer + 2, 1, AVAILABILITY_CONFIRMATION_NOTE)
    ws.cell(footer + 2, 1).alignment = Alignment(wrap_text=True)
    excluded = len(all_rows) - len(rows)
    signature_row = footer + 3
    if excluded:
        ws.merge_cells(start_row=footer + 3, start_column=1, end_row=footer + 3, end_column=6)
        ws.cell(footer + 3, 1, f"Частичный счёт: не включено неподтверждённых позиций — {excluded}.")
        signature_row += 1
    ws.merge_cells(start_row=signature_row, start_column=1, end_row=signature_row, end_column=3)
    ws.cell(signature_row, 1, f"Руководитель: __________________ / {seller['director']} /")
    ws.freeze_panes = "A12"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    if product_links:
        from openpyxl.drawing.image import Image
        products = wb.create_sheet('Товары и QR')
        products.column_dimensions['A'].width = 78
        products.column_dimensions['B'].width = 22
        products.append(['Товар — фото и характеристики', 'Открыть с телефона'])
        products['A1'].font = products['B1'].font = bold
        products.append(['На листе «Счёт» нажмите название товара. Для распечатки используйте QR-код.'])
        products['A2'].alignment = Alignment(wrap_text=True, vertical='center')
        products.row_dimensions[2].height = 32
        for row, (name, url) in enumerate(dict.fromkeys(product_links), 3):
            cell = products.cell(row, 1, name + '\n' + url)
            cell.hyperlink = url
            cell.font = Font(color='17569B', underline='single')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            qr = Image(BytesIO(product_qr_png(url)))
            qr.width = qr.height = 124
            products.add_image(qr, f'B{row}')
            products.row_dimensions[row].height = 98
        products.freeze_panes = 'A3'
        products.page_setup.orientation = 'landscape'
        products.page_setup.fitToWidth = 1
        products.page_setup.fitToHeight = 0
        products.sheet_properties.pageSetUpPr.fitToPage = True
        products.print_title_rows = '1:2'

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
