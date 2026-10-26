import copy
import html
import json
import os
import re
import tempfile
import unittest
from unittest.mock import Mock, patch

from fastapi import FastAPI
from fastapi.testclient import TestClient

import automation_pipeline as pipeline
import security
import supplier_comparison as comparison


def quote(supplier="А", name="Труба 89х4 мм", **kw):
    return dict(supplier=supplier, name=name, sku=supplier + name, unit="м", price=100, **kw)


def result_with(*offers, name="Труба 89х4 мм", quantity=60, requested_unit="м"):
    return dict(order_id=123, order_name="Трубы", markup=.05, items=[dict(position=1, requested_name=name,
        quantity=quantity, unit=requested_unit, selected=offers[0] if offers else None,
        candidates=list(offers), decision="auto_ready", proposed_unit_price=105)])


def accept_name(result, key):
    item = result["items"][0]
    candidate = next(c for c in comparison.candidates(item) if comparison.quote_key(c) == key)
    for f in comparison.offer_view(item, candidate)["features"]:
        if f["state"] == "pending":
            comparison.mutate(result, dict(action="feature", position=1, feature=f["id"], decision="accepted"))


class ComparisonLogicTests(unittest.TestCase):
    def test_shared_quantity_decision_applies_only_to_same_amount_unit_and_row(self):
        a, b, c = quote("А", offered_quantity=56), quote("Б", offered_quantity=56), quote("В", offered_quantity=55)
        result = result_with(a, b, c)
        result["items"].append({**copy.deepcopy(result["items"][0]), "position": 2})
        comparison.mutate(result, dict(action="quantity", position=1, key=comparison.quote_key(a), decision="accepted", scope="row_value"))
        row1, row2 = comparison.model(result)["rows"]
        self.assertEqual([o["quantity_state"] for o in row1["offers"]], ["accepted", "accepted", "pending"])
        self.assertTrue(all(o["quantity_state"] == "pending" for o in row2["offers"]))
        self.assertEqual(row1["quantity"], 60)
        self.assertEqual(row1["offers"][1]["shortage"], 4)
        comparison.mutate(result, dict(action="quantity", position=1, key=comparison.quote_key(b), decision="rejected", scope="row_value"))
        self.assertTrue(all(not o["eligible"] for o in comparison.model(result)["rows"][0]["offers"]))

    def test_shared_unit_acknowledgement_never_invents_a_conversion(self):
        a, b = {**quote("А"), "unit": "т", "offered_quantity": 2}, {**quote("Б"), "unit": "т", "offered_quantity": 2}
        result = result_with(a, b)
        comparison.mutate(result, dict(action="unit", position=1, key=comparison.quote_key(a), decision="accepted"))
        views = comparison.model(result)["rows"][0]["offers"]
        self.assertTrue(all(o["unit_state"] == "accepted" for o in views))
        self.assertTrue(all(o["factor"] is None and not o["eligible"] for o in views))
        comparison.mutate(result, dict(action="unit", position=1, key=comparison.quote_key(b), decision="rejected"))
        self.assertTrue(all("Единица предложения отклонена" in o["reasons"] for o in comparison.model(result)["rows"][0]["offers"]))

    def test_row_characteristic_decision_does_not_accept_short_quantity_or_other_rows(self):
        a = quote("А", "Труба 89х2 мм", offered_quantity=56)
        b = quote("Б", "Труба 89х2 мм", offered_quantity=60)
        result = result_with(a, b)
        result["items"].append({**copy.deepcopy(result["items"][0]), "position": 2})
        comparison.mutate(result, dict(action="row_features", position=1, decision="accepted"))
        row1, row2 = comparison.model(result)["rows"]
        self.assertFalse(row1["offers"][0]["eligible"])
        self.assertEqual(row1["offers"][0]["quantity_state"], "pending")
        self.assertTrue(row1["offers"][1]["eligible"])
        self.assertFalse(row2["offers"][1]["eligible"])
        comparison.mutate(result, dict(action="row_features", position=1, decision="rejected"))
        self.assertFalse(any(o["eligible"] for o in comparison.model(result)["rows"][0]["offers"]))

    def test_reject_parameter_propagates_across_suppliers_but_not_other_parameters_or_positions(self):
        a, b, c = quote("А", "Труба профильная 40х40х2 мм"), quote("Б", "Труба профильная 40x40x2,0 мм"), quote("В", "Труба профильная 40х40х3 мм")
        result = result_with(a,b,c,name="Труба профильная 40х40х3 мм")
        result["items"].append({**copy.deepcopy(result["items"][0]), "position":2})
        feature = next(f for f in comparison.offer_view(result["items"][0],a)["features"] if f["field"] == "Толщина стенки")
        comparison.mutate(result, dict(action="feature",position=1,feature=feature["id"],decision="rejected"))
        row1,row2=comparison.model(result)["rows"]
        for o in row1["offers"][:2]:
            self.assertFalse(o["eligible"])
            self.assertTrue(any(f["id"]==feature["id"] and f["state"]=="rejected" for f in o["features"]))
        self.assertTrue(row1["offers"][2]["eligible"])
        self.assertFalse(any(f["state"]=="rejected" for o in row2["offers"] for f in o["features"]))

    def test_short_quantity_requires_explicit_acceptance_and_invoice_retains_original_demand(self):
        a,b=quote("А",offered_quantity=56),quote("Б",offered_quantity=60)
        result=result_with(a,b)
        item=result["items"][0]
        self.assertFalse(comparison.offer_view(item,a)["eligible"])
        self.assertTrue(comparison.offer_view(item,b)["eligible"])
        key=comparison.quote_key(a)
        with self.assertRaises(ValueError): comparison.mutate(result,dict(action="select",position=1,key=key))
        comparison.mutate(result,dict(action="quantity",position=1,key=key,decision="accepted"))
        comparison.mutate(result,dict(action="select",position=1,key=key))
        self.assertEqual(item["quantity"],56)
        self.assertEqual(item["purchase_units"],56)
        self.assertEqual(item["comparison_source"]["quantity"],60)
        view=comparison.model(result)["rows"][0]["offers"][0]
        self.assertEqual(view["shortage"],4)
        self.assertEqual(view["total"],5600)
        self.assertTrue(pipeline._included(item))
        comparison.mutate(result,dict(action="quantity",position=1,key=key,decision="rejected"))
        self.assertFalse(pipeline._included(item))

    def test_quantity_decision_is_offer_specific_and_expires_on_quantity_change(self):
        a,b=quote("А",offered_quantity=56),quote("Б",offered_quantity=56)
        result=result_with(a,b)
        comparison.mutate(result,dict(action="quantity",position=1,key=comparison.quote_key(a),decision="accepted"))
        self.assertTrue(comparison.offer_view(result["items"][0],a)["eligible"])
        self.assertFalse(comparison.offer_view(result["items"][0],b)["eligible"])
        changed={**a,"offered_quantity":55}
        self.assertEqual(comparison.offer_view(result["items"][0],changed)["quantity_state"],"pending")

    def test_conversion_and_display_switch_are_independent(self):
        a={**quote(),"unit":"шт","name":"Труба 89х4 мм L=6м","offered_quantity":10,"price":600}
        result=result_with(a)
        item=result["items"][0]
        accept_name(result,comparison.quote_key(a))
        view=comparison.offer_view(item,a)
        self.assertEqual(view["factor"],6)
        self.assertEqual(view["converted_quantity"],60)
        self.assertEqual(view["total"],6000)
        self.assertEqual(view["unit_price"],100)
        before=copy.deepcopy(a)
        comparison.mutate(result,dict(action="mode",value="price"))
        after=comparison.offer_view(item,a)
        self.assertEqual(after["factor"],6)
        self.assertEqual(after["total"],6000)
        self.assertEqual(a,before)

    def test_tonnes_to_metres_needs_known_factor(self):
        a={**quote(),"unit":"т","offered_quantity":2,"price":50000}
        result=result_with(a)
        view=comparison.offer_view(result["items"][0],a)
        self.assertIsNone(view["factor"])
        self.assertFalse(view["eligible"])
        comparison.mutate(result,dict(action="conversion",position=1,key=view["key"],factor=30))
        view=comparison.offer_view(result["items"][0],a)
        self.assertEqual(view["converted_quantity"],60)
        self.assertAlmostEqual(view["unit_price"],50000/30,places=5)
        self.assertEqual(view["total"],100000)
        self.assertTrue(view["eligible"])

    def test_accepting_dimension_resolves_all_equivalent_spellings_without_duplicate_word_decisions(self):
        a,b=quote("А",name="Труба 89х2 мм"),quote("Б",name="Труба 89x2,0 мм")
        result=result_with(a,b)
        feature=next(f for f in comparison.offer_view(result["items"][0],a)["features"] if f["field"]=="Толщина стенки")
        comparison.mutate(result,dict(action="feature",position=1,feature=feature["id"],decision="accepted"))
        self.assertTrue(all(o["eligible"] for o in comparison.model(result)["rows"][0]["offers"]))

    def test_invoice_uses_accepted_partial_quantity_and_blocks_rejected_selection(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from invoice_generator import build_invoice_xlsx
        a=quote(offered_quantity=56)
        result=result_with(a)
        result.update(invoice_number=262,customer={"name":"ООО Тест","inn":"7716997861","kpp":"770801001"})
        key=comparison.quote_key(a)
        comparison.mutate(result,dict(action="quantity",position=1,key=key,decision="accepted"))
        comparison.mutate(result,dict(action="select",position=1,key=key))
        book=load_workbook(BytesIO(build_invoice_xlsx(result)))
        line=next(row for row in book.active.iter_rows() if row[1].value==a["name"])
        self.assertEqual(line[2].value,56)
        comparison.mutate(result,dict(action="cell",position=1,key=key,enabled=False))
        # Even a stale caller restoring the old decision cannot bypass rules.
        result["items"][0]["decision"]="approved"
        with self.assertRaisesRegex(ValueError,"подтверждённой"):
            build_invoice_xlsx(result)

    def test_pack_rounding_has_quantity_decision_and_price_basis_is_explicit(self):
        a={**quote(),"unit":"упаковка","pack_quantity":10,"offered_quantity":2,"price":100}
        result=result_with(a,quantity=12,requested_unit="шт")
        view=comparison.offer_view(result["items"][0],a)
        self.assertEqual(view["factor"],10)
        self.assertEqual(view["unit_price"],10)
        self.assertEqual(view["converted_quantity"],20)
        self.assertFalse(view["eligible"])
        self.assertTrue(view["quantity_differs"])

    def test_live_stock_shortage_is_reviewable_and_historical_stock_is_not_current(self):
        a=quote(stock=56)
        result=result_with(a)
        view=comparison.offer_view(result["items"][0],a)
        self.assertEqual(view["quantity"],56)
        self.assertFalse(view["eligible"])
        old={**a,"source_type":"price_list","source_key":"1"}
        view=comparison.offer_view(result["items"][0],old)
        self.assertEqual(view["quantity"],60)
        self.assertIsNone(view["stock"])

    def test_shared_rule_also_applies_to_new_offers_after_research(self):
        a=quote(name="Труба 89х2 мм")
        result=result_with(a)
        feature=next(f for f in comparison.offer_view(result["items"][0],a)["features"] if f["field"]=="Толщина стенки")
        comparison.mutate(result,dict(action="feature",position=1,feature=feature["id"],decision="rejected"))
        new=result_with(quote("Новый",name="Труба 89х2,0 мм"))
        comparison.preserve_comparison(new,result)
        view=comparison.model(new)["rows"][0]["offers"][0]
        self.assertFalse(view["eligible"])
        self.assertTrue(any(f["id"]==feature["id"] and f["state"]=="rejected" for f in view["features"]))
        self.assertFalse(pipeline._included(new["items"][0]))

    def test_supplier_off_and_cell_off_remove_commercial_selection(self):
        a=quote()
        result=result_with(a)
        comparison.mutate(result,dict(action="select",position=1,key=comparison.quote_key(a)))
        self.assertTrue(pipeline._included(result["items"][0]))
        comparison.mutate(result,dict(action="supplier",supplier="А",enabled=False))
        self.assertFalse(pipeline._included(result["items"][0]))
        self.assertFalse(comparison.model(result)["rows"][0]["offers"][0]["eligible"])

    def test_turning_off_other_supplier_keeps_existing_compatible_selection(self):
        a,b=quote("А"),quote("Б")
        result=result_with(a,b)
        result["items"][0].update(purchase_units=60,pack_size=1)
        comparison.mutate(result,dict(action="supplier",supplier="Б",enabled=False))
        self.assertTrue(pipeline._included(result["items"][0]))
        self.assertTrue(comparison.model(result)["rows"][0]["offers"][0]["selected"])

    def test_missing_parameters_are_reviewable_and_unknown_is_not_a_match(self):
        a=quote(name="Труба 89х4 мм")
        result=result_with(a,name="Труба 89х4 мм оцинкованная")
        missing=next(f for f in comparison.offer_view(result["items"][0],a)["features"] if f["field"]=="Покрытие")
        self.assertEqual(missing["value"],"Не указано")
        self.assertEqual(missing["state"],"pending")

    def test_invalid_or_zero_price_and_quantity_never_become_eligible(self):
        for field,value in [("price",0),("price",-1),("price","nan"),("offered_quantity",0)]:
            a={**quote(),field:value}
            result=result_with(a)
            self.assertFalse(comparison.offer_view(result["items"][0],a)["eligible"])

    def test_arbitrary_word_rule_matches_whole_word_only(self):
        a=quote(name="Болт усиленный" )
        result=result_with(a,name="Болт")
        feature=next(f for f in comparison.offer_view(result["items"][0],a)["features"] if f["value"]=="усиленный")
        comparison.mutate(result,dict(action="feature",position=1,feature=feature["id"],decision="rejected"))
        b=quote("Б",name="Болт сверхусиленный")
        self.assertFalse(any(f["state"]=="rejected" for f in comparison.offer_view(result["items"][0],b)["features"]))

    def test_sent_result_and_nonpositive_conversion_cannot_be_mutated(self):
        result=result_with(quote())
        result["live_offer_created"]=True
        with self.assertRaises(ValueError): comparison.mutate(result,dict(action="mode",value="price"))
        result["live_offer_created"]=False
        for factor in [0,-2,"NaN",float("inf")]:
            with self.assertRaises(ValueError): comparison.mutate(result,dict(action="conversion",position=1,key=comparison.quote_key(quote()),factor=factor))


class ComparisonRouteTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.patch=patch.multiple(pipeline,DB_PATH=os.path.join(self.tmp.name,"test.db"),DATABASE_URL="",_initialized_database_key=None)
        self.patch.start()
        self.app=FastAPI()
        self.fetch=Mock(side_effect=AssertionError("Comparison must use saved data"))
        pipeline.install_automation_pipeline(self.app,self.fetch)
        comparison.install_supplier_comparison(self.app)
        self.session=patch.object(security,"_valid_session",return_value=True)
        self.session.start()
        self.client=TestClient(self.app)
        self.a=quote(offered_quantity=56)
        self.result=result_with(self.a)
        with pipeline._connect() as conn:
            cursor=conn.execute("INSERT INTO automation_jobs (dedupe_key,order_id,event_type,status,result_json,created_at,updated_at) VALUES ('comparison',123,'manual_order','needs_review',?,'now','now')",(json.dumps(self.result),))
            self.job=cursor.lastrowid
        self.url=f"/dashboard/automation/jobs/{self.job}/comparison"

    def tearDown(self):
        self.client.close();self.session.stop();self.patch.stop();self.tmp.cleanup()

    def data(self):
        response=self.client.get(self.url)
        self.assertEqual(response.status_code,200)
        return json.loads(re.search(r'id="comparison-data">(.*?)</script>',response.text,re.S)[1])

    def post(self,**action):
        data=self.data()
        return self.client.post(self.url+"/action",json={"csrf":data["csrf"],"revision":data["revision"],**action})

    def test_get_is_read_only_and_scripts_are_private_with_strict_csp(self):
        before=self.data()
        self.assertEqual(before["rows"][0]["offers"][0]["quantity_state"],"pending")
        after=self.data()
        self.assertEqual(before["revision"],after["revision"])
        self.fetch.assert_not_called()
        self.assertEqual(self.client.get('/dashboard/comparison/assets/comparison.js').status_code,200)
        self.assertEqual(self.client.get('/dashboard/comparison/assets/secrets').status_code,404)
        self.assertIn("script-src 'nonce-",self.client.get(self.url).headers["content-security-policy"])
        with patch.object(security,"_valid_session",return_value=False):
            self.assertEqual(self.client.get(self.url).status_code,401)

    def test_original_matrix_state_persists_without_changing_review(self):
        data=self.data();matrix=data['prototype'];state=matrix['state']
        state['cellOn']={'1_0':False};state['diffState']={'1_0_qty':'accepted'}
        response=self.post(action='prototype_state',signature=matrix['signature'],state=state)
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(self.data()['prototype']['state'],state)
        with pipeline._connect() as conn:
            row=conn.execute('SELECT result_json,status FROM automation_jobs WHERE id=?',(self.job,)).fetchone()
        result=json.loads(row['result_json']);result.pop('prototype_comparison')
        self.assertEqual(result,self.result)
        self.assertEqual(row['status'],'needs_review')

    def test_sent_invoice_allows_matrix_decisions_without_mutating_invoice(self):
        self.result['live_offer_created'] = True
        self.result['invoice_number'] = 278
        with pipeline._connect() as conn:
            conn.execute("UPDATE automation_jobs SET status='offer_created',result_json=?,invoice_number=278 WHERE id=?",(json.dumps(self.result),self.job))
        data=self.data()
        self.assertFalse(data['read_only'])
        matrix=data['prototype'];state=matrix['state']
        state['diffState']={'1_0_qty':'accepted'}
        response=self.post(action='prototype_state',signature=matrix['signature'],state=state)
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(self.data()['prototype']['state'],state)
        with pipeline._connect() as conn:
            row=conn.execute('SELECT result_json,status,invoice_number FROM automation_jobs WHERE id=?',(self.job,)).fetchone()
        saved=json.loads(row['result_json']);saved.pop('prototype_comparison')
        self.assertEqual(saved,self.result)
        self.assertEqual((row['status'],row['invoice_number']),('offer_created',278))
        self.assertEqual(self.post(action='mode',value='price').status_code,409)

    def test_processing_matrix_remains_locked(self):
        with pipeline._connect() as conn:
            conn.execute("UPDATE automation_jobs SET status='processing' WHERE id=?",(self.job,))
        data=self.data();self.assertTrue(data['read_only'])
        matrix=data['prototype']
        self.assertEqual(self.post(action='prototype_state',signature=matrix['signature'],state=matrix['state']).status_code,409)

    def test_quantity_choice_is_saved_and_survives_page_reload(self):
        key=comparison.quote_key(self.a)
        response=self.post(action="quantity",position=1,key=key,decision="accepted")
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(self.data()["rows"][0]["offers"][0]["quantity_state"],"accepted")
        response=self.post(action="select",position=1,key=key)
        self.assertEqual(response.status_code,200,response.text)
        with pipeline._connect() as conn:
            result=json.loads(conn.execute("SELECT result_json FROM automation_jobs WHERE id=?",(self.job,)).fetchone()[0])
        self.assertEqual(result["items"][0]["quantity"],56)
        self.assertTrue(pipeline._included(result["items"][0]))
        self.assertIsNotNone(result.get("invoice_number"))

    def test_normal_review_save_preserves_conversion_and_cannot_bypass_quantity_decision(self):
        key=comparison.quote_key(self.a)
        self.assertEqual(self.post(action="quantity",position=1,key=key,decision="accepted").status_code,200)
        self.assertEqual(self.post(action="select",position=1,key=key).status_code,200)
        review=f"/dashboard/automation/jobs/{self.job}/review"
        response=self.client.post(review,data={"include_1":"1","candidate_1":"0","quantity_1":"56","price_1":"110"},follow_redirects=False)
        self.assertEqual(response.status_code,303)
        self.assertTrue(self.data()["rows"][0]["offers"][0]["selected"])
        response=self.client.post(review,data={"include_1":"1","candidate_1":"0","quantity_1":"60","price_1":"110"},follow_redirects=False)
        self.assertEqual(response.status_code,409)

    def test_csrf_and_stale_revision_cannot_write(self):
        data=self.data()
        action=dict(action="mode",value="price",revision=data["revision"])
        self.assertEqual(self.client.post(self.url+"/action",json=action).status_code,403)
        self.assertEqual(self.post(action="mode",value="price").status_code,200)
        self.assertEqual(self.client.post(self.url+"/action",json={**action,"csrf":data["csrf"]}).status_code,409)

    def test_active_submission_and_sent_jobs_are_immutable(self):
        with pipeline._connect() as conn:
            conn.execute("UPDATE automation_jobs SET status='processing' WHERE id=?",(self.job,))
        self.assertEqual(self.post(action="mode",value="price").status_code,409)

    def test_embedded_names_cannot_close_json_script_or_create_script_elements(self):
        with pipeline._connect() as conn:
            result=copy.deepcopy(self.result);result["order_name"]="</script><script>alert(1)</script>"
            conn.execute("UPDATE automation_jobs SET result_json=? WHERE id=?",(json.dumps(result),self.job))
        response=self.client.get(self.url)
        self.assertNotIn("<script>alert(1)</script>",response.text)
        self.assertEqual(self.data()["order_name"],"</script><script>alert(1)</script>")


if __name__ == '__main__':
    unittest.main()
