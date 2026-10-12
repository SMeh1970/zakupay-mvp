"""Durable supplier price imports and indexed candidate retrieval.

Original files and versions live in the same database as requests. Only an
explicit, validated import activates prices; request/invoice snapshots stay put.
"""
from __future__ import annotations

import csv
import hashlib
import html
import io
import json
import math
import re
import zipfile
from contextlib import contextmanager
from datetime import date, datetime, timezone
from pathlib import Path
from urllib.parse import quote, urlparse

from fastapi import HTTPException, Request
from fastapi.responses import HTMLResponse, RedirectResponse, Response
from starlette.concurrency import run_in_threadpool

MAX_FILE_BYTES = 12 * 1024 * 1024
MAX_ROWS = 100000
MAX_COLUMNS = 200
FIELDS = {"name": "Наименование", "price": "Закупочная цена", "article": "Артикул",
          "brand": "Бренд", "unit": "Единица цены", "stock": "Остаток на дату прайса",
          "pack_size": "Штук в упаковке", "url": "Ссылка на товар"}
ALIASES = {
    "name": {"наименование", "наименование товара", "товар", "название", "номенклатура", "name", "product name"},
    "price": {"цена", "цена руб", "цена с ндс", "цена без ндс", "закупочная цена", "оптовая цена", "price"},
    "article": {"артикул", "код", "код товара", "sku", "article"},
    "brand": {"бренд", "производитель", "торговая марка", "brand"},
    "unit": {"ед", "ед изм", "единица", "единица измерения", "unit"},
    "stock": {"остаток", "наличие", "остаток на складе", "stock"},
    "pack_size": {"штук в упаковке", "количество в упаковке", "кратность упаковки", "pack size"},
    "url": {"ссылка", "ссылка на товар", "url"},
}


def norm(value):
    return " ".join(re.findall(r"[a-zа-я0-9]+", str(value or "").lower().replace("ё", "е")))


def tokens(value):
    stop = {"для", "шт", "штук", "мм", "см", "руб", "набор", "комплект", "упаковка"}
    result = set()
    for token in norm(value).split():
        if len(token) < 2 or token in stop:
            continue
        # A small shared index stem handles common catalogue inflections; the
        # existing full matcher still checks brands, types, articles and sizes.
        if re.fullmatch(r"[а-я]{5,}", token):
            token = re.sub(r"(?:ами|ями|ого|ему|ому|ыми|ими|ый|ий|ая|ое|ые|ие|ой|ей|ов|ев|ам|ям|ах|ях|ы|и|а|я|у|ю|е)$", "", token)
        result.add(token)
    return sorted(result)[:32]


def number(value):
    if isinstance(value, bool) or value is None:
        return None
    text = str(value).strip().replace("\xa0", "").replace("\u202f", "").replace(" ", "")
    text = re.sub(r"(?:₽|руб\.?|RUB)$", "", text, flags=re.I)
    if "," in text and "." in text:
        text = text.replace(",", "") if text.rfind(".") > text.rfind(",") else text.replace(".", "").replace(",", ".")
    else:
        text = text.replace(",", ".")
    try:
        result = float(text)
        return result if math.isfinite(result) else None
    except (ValueError, TypeError):
        return None


def safe_url(value):
    value = str(value or "").strip()
    parsed = urlparse(value)
    return value if parsed.scheme in {"http", "https"} and parsed.netloc else ""


def initialize_schema(conn, postgres):
    serial = "BIGSERIAL" if postgres else "INTEGER"
    blob = "BYTEA" if postgres else "BLOB"
    auto = "" if postgres else " AUTOINCREMENT"
    conn.execute(f"""CREATE TABLE IF NOT EXISTS supplier_price_lists (
        id {serial} PRIMARY KEY{auto}, catalog_key TEXT NOT NULL, supplier TEXT NOT NULL,
        catalog_name TEXT NOT NULL, filename TEXT NOT NULL, file_bytes {blob} NOT NULL,
        content_hash TEXT NOT NULL, status TEXT NOT NULL, price_date TEXT NOT NULL,
        source_url TEXT, settings_json TEXT, imported_rows INTEGER NOT NULL DEFAULT 0,
        skipped_rows INTEGER NOT NULL DEFAULT 0, created_at TEXT NOT NULL, activated_at TEXT
    )""")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_supplier_price_lists_active ON supplier_price_lists(status,catalog_key)")
    conn.execute(f"""CREATE TABLE IF NOT EXISTS supplier_price_items (
        id {serial} PRIMARY KEY{auto}, list_id BIGINT NOT NULL REFERENCES supplier_price_lists(id),
        name_norm TEXT NOT NULL, article_norm TEXT NOT NULL, price DOUBLE PRECISION NOT NULL,
        payload_json TEXT NOT NULL
    )""")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_supplier_price_items_list ON supplier_price_items(list_id)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_supplier_price_items_article ON supplier_price_items(article_norm)")
    conn.execute(f"""CREATE TABLE IF NOT EXISTS supplier_price_tokens (
        item_id BIGINT NOT NULL REFERENCES supplier_price_items(id), token TEXT NOT NULL,
        PRIMARY KEY(token,item_id)
    )""")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_supplier_price_tokens_item ON supplier_price_tokens(item_id)")
    # Concurrent imports of the same catalogue cannot leave two active versions.
    conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_supplier_price_active_version ON supplier_price_lists(catalog_key) WHERE status='active'")


@contextmanager
def database():
    from automation_pipeline import _connect
    conn = _connect()
    try:
        with conn:
            yield conn
    finally:
        conn.close()


def execute(conn, sql, params=()):
    from automation_pipeline import _execute
    return _execute(conn, sql, params)


def table_rows(content, filename, sheet="", limit=MAX_ROWS + 1):
    """Return one sheet, retaining physical row numbers and cached formula values."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        for encoding in ("utf-8-sig", "cp1251"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError("Не удалось прочитать кодировку CSV")
        try:
            dialect = csv.Sniffer().sniff(text[:16000], delimiters=";,\t")
        except csv.Error:
            dialect = csv.excel
        rows = []
        for row in csv.reader(io.StringIO(text), dialect):
            if len(row) > MAX_COLUMNS:
                raise ValueError("В таблице больше 200 столбцов")
            rows.append(row)
            if len(rows) >= limit:
                break
        return ["CSV"], "CSV", rows
    if suffix == ".xls":
        import xlrd
        book = xlrd.open_workbook(file_contents=content, on_demand=True)
        try:
            names = book.sheet_names()
            selected = sheet if sheet in names else names[0]
            page = book.sheet_by_name(selected)
            if page.ncols > MAX_COLUMNS:
                raise ValueError("В таблице больше 200 столбцов")
            return names, selected, [page.row_values(i) for i in range(min(page.nrows, limit))]
        finally:
            book.release_resources()
    if suffix != ".xlsx":
        raise ValueError("Поддерживаются Excel XLSX, XLS и CSV")
    from openpyxl import load_workbook
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        if sum(entry.file_size for entry in archive.infolist()) > 150 * 1024 * 1024:
            raise ValueError("Слишком большой распакованный Excel-файл")
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        names = book.sheetnames
        selected = sheet if sheet in names else names[0]
        page = book[selected]
        if page.max_column and page.max_column > MAX_COLUMNS:
            raise ValueError("В таблице больше 200 столбцов. Удалите пустые столбцы с форматированием.")
        rows = []
        for row in page.iter_rows(values_only=True):
            rows.append(list(row))
            if len(rows) >= limit:
                break
        return names, selected, rows
    finally:
        book.close()


def detect_header(rows):
    best, score = 1, -1
    for index, row in enumerate(rows[:40], 1):
        values = {norm(value) for value in row if value is not None}
        current = sum(bool(values & aliases) for aliases in ALIASES.values())
        if current > score:
            best, score = index, current
    return best


def detect_mapping(header):
    mapping = {}
    for field, aliases in ALIASES.items():
        matches = [i for i, value in enumerate(header) if norm(value) in aliases]
        # Multiple price columns require an explicit operator choice.
        mapping[field] = matches[0] if len(matches) == 1 else None
    return mapping


def get_list(list_id, include_file=False):
    columns = "*" if include_file else "id,catalog_key,supplier,catalog_name,filename,status,price_date,source_url,settings_json,imported_rows,skipped_rows,created_at,activated_at"
    with database() as conn:
        row = execute(conn, f"SELECT {columns} FROM supplier_price_lists WHERE id=?", (list_id,)).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="Прайс не найден")
    return dict(row)


def stage_file(content, filename, supplier, catalog_name, price_date, source_url=""):
    supplier, catalog_name = supplier.strip(), catalog_name.strip() or "Основной"
    if not supplier or len(supplier) > 200 or len(catalog_name) > 200:
        raise ValueError("Укажите поставщика и название прайса длиной до 200 символов")
    parsed_date = date.fromisoformat(price_date)
    if parsed_date > date.today():
        raise ValueError("Дата прайса не может быть в будущем")
    if not content or len(content) > MAX_FILE_BYTES:
        raise ValueError("Выберите файл размером до 12 МБ")
    filename = Path(filename.replace("\\", "/")).name[:200]
    _, _, rows = table_rows(content, filename, limit=45)
    if not rows:
        raise ValueError("В файле нет строк")
    now = datetime.now(timezone.utc).isoformat()
    key = hashlib.sha256((norm(supplier) + "|" + norm(catalog_name)).encode()).hexdigest()
    with database() as conn:
        row = execute(conn, """INSERT INTO supplier_price_lists
            (catalog_key,supplier,catalog_name,filename,file_bytes,content_hash,status,price_date,source_url,created_at)
            VALUES (?,?,?,?,?,?,'pending',?,?,?) RETURNING id""",
            (key, supplier, catalog_name, filename, content, hashlib.sha256(content).hexdigest(), price_date, safe_url(source_url), now)).fetchone()
    return row["id"]


def parse_items(record, settings):
    names, sheet, rows = table_rows(bytes(record["file_bytes"]), record["filename"], settings["sheet"])
    if settings["sheet"] not in names:
        raise ValueError("Выбранный лист отсутствует в файле")
    header_row = int(settings["header_row"])
    if not 1 <= header_row <= min(100, len(rows)):
        raise ValueError("Укажите номер строки заголовков от 1 до 100")
    width = len(rows[header_row - 1])
    mapping = settings["mapping"]
    for field in FIELDS:
        column = mapping.get(field)
        if column is not None and (not isinstance(column, int) or column < 0 or column >= width):
            raise ValueError("Некорректный столбец: " + FIELDS[field])
    if mapping.get("name") is None or mapping.get("price") is None:
        raise ValueError("Выберите столбцы наименования и закупочной цены")
    used = [value for value in mapping.values() if value is not None]
    if len(used) != len(set(used)):
        raise ValueError("Один столбец нельзя назначить нескольким полям")
    vat_mode = settings.get("vat_mode")
    if vat_mode not in {"included", "excluded", "exempt"}:
        raise ValueError("Укажите, включён ли НДС в цены")
    vat_rate = number(settings.get("vat_rate"))
    if vat_rate is None or not 0 <= vat_rate <= 100:
        raise ValueError("Ставка НДС должна быть от 0 до 100%")
    default_unit = str(settings.get("default_unit") or "").strip()
    if mapping.get("unit") is None and not default_unit:
        raise ValueError("Выберите столбец единиц или задайте единицу цены")
    if len(rows) > MAX_ROWS:
        raise ValueError("В файле больше 100 000 строк. Разделите прайс на несколько файлов.")
    items, errors, skipped = [], [], 0
    for row_number, row in enumerate(rows[header_row:], header_row + 1):
        if not any(value not in (None, "") for value in row):
            continue
        def cell(field):
            col = mapping.get(field)
            return row[col] if col is not None and col < len(row) else None
        name = str(cell("name") or "").strip()
        raw_price = cell("price")
        # Rubric rows/notes with no price are reported as skipped; populated
        # malformed product prices block activation instead of losing goods.
        if not name or raw_price in (None, ""):
            skipped += 1
            continue
        price = number(raw_price)
        unit = str(cell("unit") or default_unit).strip()
        pack = number(cell("pack_size"))
        stock = number(cell("stock"))
        if price is None or price <= 0 or not unit or (pack is not None and (pack < 1 or not pack.is_integer())):
            errors.append(f"Строка {row_number}: проверьте цену, единицу и размер упаковки")
            continue
        if len(name) > 4000:
            errors.append(f"Строка {row_number}: слишком длинное наименование")
            continue
        gross_price = round(price * (1 + vat_rate / 100), 6) if vat_mode == "excluded" else price
        article = str(cell("article") or "").strip()
        if isinstance(cell("article"), float) and cell("article").is_integer():
            article = str(int(cell("article")))
        identity = article or f"price:{record['catalog_key']}:{sheet}:{row_number}"
        items.append({"supplier": record["supplier"], "name": name, "article": article or None,
            "sku": identity, "brand": str(cell("brand") or "").strip() or None,
            "unit": unit, "price": gross_price, "stock": None,
            "url": safe_url(cell("url")) or None, "price_type": "Прайс поставщика",
            "source_type": "price_list", "source_list_id": record["id"],
            "source_key": record["catalog_key"], "source_filename": record["filename"],
            "source_sheet": sheet, "source_row": row_number, "price_date": record["price_date"],
            "source_url": record.get("source_url") or None, "original_price": price,
            "vat_mode": vat_mode, "vat_rate": vat_rate, "pack_quantity": int(pack) if pack else None,
            "price_list_stock": stock if stock is not None and stock >= 0 else None})
    if errors:
        raise ValueError("; ".join(errors[:8]) + (f"; всего ошибок: {len(errors)}" if len(errors) > 8 else ""))
    if not items:
        raise ValueError("Не найдено ни одной позиции с положительной ценой. Проверьте лист и столбцы.")
    return items, skipped


def activate_list(list_id, settings):
    record = get_list(list_id, include_file=True)
    settings_json = json.dumps(settings, ensure_ascii=False, sort_keys=True)
    items, skipped = parse_items(record, settings)
    from automation_pipeline import DATABASE_URL, _lock
    with _lock, database() as conn:
        # Serialise catalogue replacement across web/worker processes as well.
        if DATABASE_URL:
            execute(conn, "SELECT pg_advisory_xact_lock(?)", (int(record["catalog_key"][:15], 16),))
        current = execute(conn, "SELECT status FROM supplier_price_lists WHERE id=?", (list_id,)).fetchone()
        if current["status"] != "pending":
            raise ValueError("Этот прайс уже обработан. Для обновления загрузите новый файл.")
        duplicate = execute(conn, """SELECT id FROM supplier_price_lists
            WHERE catalog_key=? AND status='active' AND content_hash=? AND settings_json=? AND price_date=?""",
            (record["catalog_key"], record["content_hash"], settings_json, record["price_date"])).fetchone()
        if duplicate:
            execute(conn, "UPDATE supplier_price_lists SET status='duplicate' WHERE id=?", (list_id,))
            return {"id": duplicate["id"], "duplicate": True, "imported": len(items), "skipped": skipped}
        index_rows = []
        for offset in range(0, len(items), 100):
            batch = items[offset:offset + 100]
            params = []
            for item in batch:
                params.extend((list_id, norm(item["name"]), re.sub(r"[^a-zа-я0-9]", "", norm(item["article"])),
                               item["price"], json.dumps(item, ensure_ascii=False)))
            inserted = execute(conn, "INSERT INTO supplier_price_items (list_id,name_norm,article_norm,price,payload_json) VALUES "
                + ",".join("(?,?,?,?,?)" for _ in batch) + " RETURNING id,payload_json", params).fetchall()
            for row in inserted:
                item = json.loads(row["payload_json"])
                index_rows.extend((row["id"], token) for token in tokens(" ".join(str(item.get(key) or "") for key in ("name", "article", "brand"))))
        cursor = conn.cursor()
        try:
            cursor.executemany("INSERT INTO supplier_price_tokens(item_id,token) VALUES (%s,%s)" if DATABASE_URL else
                               "INSERT INTO supplier_price_tokens(item_id,token) VALUES (?,?)", index_rows)
        finally:
            cursor.close()
        execute(conn, "UPDATE supplier_price_lists SET status='superseded' WHERE catalog_key=? AND status='active'", (record["catalog_key"],))
        execute(conn, """UPDATE supplier_price_lists SET status='active',settings_json=?,imported_rows=?,
            skipped_rows=?,activated_at=? WHERE id=?""",
            (settings_json, len(items), skipped, datetime.now(timezone.utc).isoformat(), list_id))
    return {"id": list_id, "duplicate": False, "imported": len(items), "skipped": skipped}


def search_prices(query, limit=8):
    from supplier_adapters import SupplierQuote
    from vi_order_match import _identifiers, _score_details
    from automation_pipeline import _hard_conflicts, _candidate_matches_brand, _requested_brands
    terms = tokens(query)
    ids = sorted(_identifiers(query))
    if not terms and not ids:
        return []
    rows = {}
    with database() as conn:
        if ids:
            placeholders = ",".join("?" for _ in ids)
            exact = execute(conn, f"""SELECT i.id,i.payload_json FROM supplier_price_items i
                JOIN supplier_price_lists p ON p.id=i.list_id
                WHERE p.status='active' AND i.article_norm IN ({placeholders}) ORDER BY i.price LIMIT 300""",
                [identifier.lower() for identifier in ids]).fetchall()
            rows.update((row["id"], row) for row in exact)
        if terms:
            placeholders = ",".join("?" for _ in terms)
            matched = execute(conn, f"""SELECT i.id,i.payload_json,COUNT(*) AS hits FROM supplier_price_tokens t
                JOIN supplier_price_items i ON i.id=t.item_id JOIN supplier_price_lists p ON p.id=i.list_id
                WHERE p.status='active' AND t.token IN ({placeholders})
                GROUP BY i.id,i.payload_json,i.price ORDER BY hits DESC,i.price ASC LIMIT 400""", terms).fetchall()
            rows.update((row["id"], row) for row in matched)
    quotes = [SupplierQuote(**json.loads(row["payload_json"])) for row in rows.values()]
    brands = _requested_brands(query, [item.brand for item in quotes])
    accepted = []
    for item in quotes:
        score, _ = _score_details(query, item)
        if score < 0.48 or _hard_conflicts(query, item.to_dict()) or (brands and not _candidate_matches_brand(item.to_dict(), brands)):
            continue
        accepted.append((item, score))
    # Return a shortlist per supplier so a large cheap catalogue cannot hide
    # alternatives from every other firm. Final unit-aware selection is shared.
    accepted.sort(key=lambda pair: (-pair[1], pair[0].price) if brands else
                  (pair[0].price / (pair[0].pack_quantity or 1), -pair[1]))
    result, counts = [], {}
    for item, _ in accepted:
        if counts.get(item.supplier, 0) >= max(8, limit):
            continue
        result.append(item)
        counts[item.supplier] = counts.get(item.supplier, 0) + 1
        if len(result) >= 80:
            break
    return result


def page(body, title="Прайсы поставщиков"):
    return HTMLResponse("<!doctype html><html lang='ru'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
        f"<title>{html.escape(title)}</title><style>body{{font-family:Arial;background:#f4f6f8;color:#202124;margin:0}}main{{max-width:1150px;margin:auto;padding:24px}}section{{background:#fff;padding:20px;border-radius:12px;margin:18px 0}}label{{display:block;margin:12px 0}}input,select{{padding:9px;max-width:100%;box-sizing:border-box}}button,.button{{background:#1a73e8;color:white;padding:11px 15px;border:0;border-radius:7px;cursor:pointer;text-decoration:none;display:inline-block}}table{{border-collapse:collapse;width:100%}}td,th{{padding:9px;border:1px solid #ddd;text-align:left}}.scroll{{overflow:auto}}.note{{color:#5f6368}}.error{{color:#b3261e}}.grid{{display:grid;grid-template-columns:repeat(auto-fit,minmax(210px,1fr));gap:12px}}a{{color:#174ea6}}</style></head>"
        "<body><main><p><a href='/dashboard/automation'>← Заявки</a> · <a href='/dashboard/prices'>Прайсы</a></p>" + body + "</main></body></html>",
        headers={"Cache-Control": "no-store"})


def preview_page(list_id, sheet="", header_row=None, error="", settings=None):
    record = get_list(list_id, include_file=True)
    sheets, selected, rows = table_rows(bytes(record["file_bytes"]), record["filename"], sheet, limit=115)
    header_row = header_row or detect_header(rows)
    if not 1 <= header_row <= min(100, len(rows)):
        raise ValueError("Некорректная строка заголовков")
    header = rows[header_row - 1]
    mapping = (settings or {}).get("mapping") or detect_mapping(header)
    esc = lambda value: html.escape(str(value or ""), quote=True)
    options = "".join(f"<option value='{esc(name)}' {'selected' if name == selected else ''}>{esc(name)}</option>" for name in sheets)
    controls = f"<form method='post' action='/dashboard/prices/{list_id}/preview'><label>Лист <select name='sheet'>{options}</select></label><label>Строка заголовков <input name='header_row' type='number' min='1' max='100' value='{header_row}'></label><button>Показать лист и столбцы</button></form>"
    selects = []
    for field, label in FIELDS.items():
        cols = "<option value=''>Не используется</option>" + "".join(
            f"<option value='{index}' {'selected' if mapping.get(field) == index else ''}>{index + 1}: {esc(value) or '(без заголовка)'}</option>" for index, value in enumerate(header))
        selects.append(f"<label>{label}{' *' if field in {'name','price'} else ''}<br><select name='column_{field}'>{cols}</select></label>")
    sample = "".join("<tr><td>" + str(index) + "</td>" + "".join(f"<td>{esc(value)}</td>" for value in row[:len(header)]) + "</tr>" for index, row in enumerate(rows[header_row:header_row + 8], header_row + 1))
    settings = settings or {}
    vat_options = "<option value=''>Выберите</option>" + "".join(f"<option value='{value}' {'selected' if settings.get('vat_mode') == value else ''}>{label}</option>" for value, label in [("included", "Цена включает НДС"), ("excluded", "Цена без НДС — добавить указанную ставку"), ("exempt", "Поставщик продаёт без НДС")])
    return page(f"<h1>Проверка прайса</h1><p>{esc(record['supplier'])} · {esc(record['filename'])} · дата цен: {esc(record['price_date'])}</p>"
        + (f"<p class='error'>{esc(error)}</p>" if error else "") + f"<section>{controls}</section><section><h2>Соответствие столбцов</h2>"
        f"<form method='post' action='/dashboard/prices/{list_id}/import'><input type='hidden' name='sheet' value='{esc(selected)}'><input type='hidden' name='header_row' value='{header_row}'>"
        f"<div class='grid'>{''.join(selects)}</div><label>Единица цены, если в строке не указана <input name='default_unit' value='{esc(settings.get('default_unit'))}' placeholder='Например: шт, м, кг, упак'></label>"
        f"<label>НДС в прайсе * <select name='vat_mode' required>{vat_options}</select></label><label>Ставка НДС, % <input name='vat_rate' type='number' min='0' max='100' step='0.01' value='{esc(settings.get('vat_rate', 22))}' required></label>"
        "<p class='note'>Выберите именно закупочную цену. Импортируется выбранный лист. Этот файл полностью заменит активную версию прайса с тем же поставщиком и названием; старые версии и исходные файлы сохраняются. Остаток из прайса требует подтверждения поставщиком.</p><button>Импортировать и использовать в поиске</button></form></section>"
        "<section class='scroll'><h2>Пример исходных строк</h2><table><tr><th>Строка</th>" + "".join(f"<th>{esc(value)}</th>" for value in header) + f"</tr>{sample}</table></section>")


def install_price_lists(app):
    @app.get("/dashboard/prices")
    def dashboard_prices(message: str = "", q: str = ""):
        with database() as conn:
            lists = execute(conn, """SELECT id,supplier,catalog_name,filename,status,price_date,imported_rows,skipped_rows
                FROM supplier_price_lists ORDER BY id DESC LIMIT 100""").fetchall()
        esc = lambda value: html.escape(str(value or ""), quote=True)
        statuses = {"active": "В поиске", "pending": "Ожидает проверки", "disabled": "Отключён", "superseded": "Заменён новой версией", "duplicate": "Повторная загрузка"}
        result_rows = []
        for item in lists:
            lid = item["id"]
            action = f"<a href='/dashboard/prices/{lid}/preview'>Проверить столбцы</a>" if item["status"] == "pending" else ""
            if item["status"] in {"active", "disabled"}:
                action = f"<form method='post' action='/dashboard/prices/{lid}/toggle'><button>{'Отключить' if item['status'] == 'active' else 'Включить'}</button></form>"
            result_rows.append(f"<tr><td>{esc(item['supplier'])}<br>{esc(item['catalog_name'])}</td><td><a href='/dashboard/prices/{lid}/file'>{esc(item['filename'])}</a></td><td>{esc(item['price_date'])}</td><td>{item['imported_rows']}<br><small>пропущено: {item['skipped_rows']}</small></td><td>{statuses.get(item['status'],esc(item['status']))}</td><td>{action}</td></tr>")
        search_html = ""
        if q.strip():
            from automation_pipeline import _purchase_label
            found = search_prices(q.strip(), 8)
            search_html = "<ul>" + "".join(f"<li>{esc(_purchase_label(item.to_dict(), ''))}</li>" for item in found) + "</ul>" if found else "<p>Подходящие товары в активных прайсах не найдены.</p>"
        return page("<h1>Прайсы поставщиков</h1><p>Импортированные цены участвуют в поиске внутри заявки вместе с API поставщиков.</p>"
            + (f"<p>{esc(message)}</p>" if message else "") + "<section><h2>Загрузить прайс</h2><form method='post' action='/dashboard/prices/upload' enctype='multipart/form-data'>"
            "<div class='grid'><label>Поставщик *<br><input name='supplier' required maxlength='200'></label><label>Название прайса<br><input name='catalog_name' value='Основной' maxlength='200'></label>"
            f"<label>Дата цен *<br><input name='price_date' type='date' value='{date.today().isoformat()}' required></label><label>Адрес первоисточника<br><input name='source_url' type='url' placeholder='https://...'></label></div>"
            "<label>Excel XLSX, XLS или CSV, до 12 МБ *<br><input name='file' type='file' accept='.xlsx,.xls,.csv' required></label><p class='note'>Импортируются цены в рублях. Для нескольких листов загрузите их как отдельные прайсы с разными названиями.</p><button>Загрузить и проверить</button></form></section>"
            f"<section><h2>Проверить поиск в прайсах</h2><form method='get'><input name='q' value='{esc(q)}' placeholder='Наименование или артикул' required><button>Найти</button></form>{search_html}</section>"
            "<section class='scroll'><h2>Последние 100 загрузок</h2><table><tr><th>Поставщик / прайс</th><th>Файл</th><th>Дата цен</th><th>Позиций</th><th>Статус</th><th>Действие</th></tr>"
            + ("".join(result_rows) or "<tr><td colspan='6'>Прайсы ещё не загружены.</td></tr>") + "</table></section>")

    @app.post("/dashboard/prices/upload")
    async def upload_price(request: Request):
        form = await request.form()
        upload = form.get("file")
        if not upload or not hasattr(upload, "read"):
            return page("<h1>Прайс не загружен</h1><p class='error'>Выберите файл.</p>")
        try:
            content = await upload.read(MAX_FILE_BYTES + 1)
            list_id = await run_in_threadpool(stage_file, content, upload.filename or "", str(form.get("supplier") or ""),
                str(form.get("catalog_name") or ""), str(form.get("price_date") or ""), str(form.get("source_url") or ""))
        except Exception as exc:
            if isinstance(exc, HTTPException):
                raise
            return page("<h1>Прайс не загружен</h1><p class='error'>" + html.escape(str(exc)) + "</p><a href='/dashboard/prices'>Вернуться к загрузке</a>")
        finally:
            await upload.close()
        return RedirectResponse(f"/dashboard/prices/{list_id}/preview", status_code=303)

    @app.get("/dashboard/prices/{list_id}/preview")
    def get_preview(list_id: int):
        return preview_page(list_id)

    @app.post("/dashboard/prices/{list_id}/preview")
    async def change_preview(list_id: int, request: Request):
        form = await request.form()
        try:
            return await run_in_threadpool(preview_page, list_id, str(form.get("sheet") or ""), int(form.get("header_row") or 1))
        except ValueError as exc:
            return page("<p class='error'>" + html.escape(str(exc)) + "</p>")

    @app.post("/dashboard/prices/{list_id}/import")
    async def import_price(list_id: int, request: Request):
        form = await request.form()
        settings = {"sheet": str(form.get("sheet") or ""), "header_row": str(form.get("header_row") or "1"),
            "mapping": {}, "default_unit": str(form.get("default_unit") or ""),
            "vat_mode": str(form.get("vat_mode") or ""), "vat_rate": str(form.get("vat_rate") or "")}
        try:
            settings["mapping"] = {field: int(form[f"column_{field}"]) if form.get(f"column_{field}") not in (None, "") else None for field in FIELDS}
            result = await run_in_threadpool(activate_list, list_id, settings)
        except ValueError as exc:
            return await run_in_threadpool(preview_page, list_id, settings["sheet"], int(settings["header_row"]), str(exc), settings)
        message = "Этот файл уже используется в поиске." if result["duplicate"] else f"Импортировано позиций: {result['imported']}. Пропущено строк без наименования или цены: {result['skipped']}."
        return RedirectResponse("/dashboard/prices?message=" + quote(message), status_code=303)

    @app.post("/dashboard/prices/{list_id}/toggle")
    def toggle_price(list_id: int):
        from automation_pipeline import DATABASE_URL, _lock
        record = get_list(list_id)
        with _lock, database() as conn:
            if DATABASE_URL:
                execute(conn, "SELECT pg_advisory_xact_lock(?)", (int(record["catalog_key"][:15], 16),))
            current = execute(conn, "SELECT status FROM supplier_price_lists WHERE id=?", (list_id,)).fetchone()["status"]
            if current == "active":
                execute(conn, "UPDATE supplier_price_lists SET status='disabled' WHERE id=?", (list_id,))
            elif current == "disabled":
                execute(conn, "UPDATE supplier_price_lists SET status='superseded' WHERE catalog_key=? AND status='active'", (record["catalog_key"],))
                execute(conn, "UPDATE supplier_price_lists SET status='active' WHERE id=?", (list_id,))
            else:
                raise HTTPException(status_code=409, detail="Можно включать и отключать только импортированный прайс")
        return RedirectResponse("/dashboard/prices", status_code=303)

    @app.get("/dashboard/prices/{list_id}/file")
    def download_original(list_id: int):
        record = get_list(list_id, include_file=True)
        return Response(bytes(record["file_bytes"]), media_type="application/octet-stream", headers={
            "Content-Disposition": "attachment; filename*=UTF-8''" + quote(record["filename"]), "Cache-Control": "no-store"})
